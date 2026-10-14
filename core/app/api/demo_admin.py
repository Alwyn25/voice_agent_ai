"""Demo admin API (docs/gtm-demo-links.md §5, §6, §16.4): industries, workbook imports,
company records, facts, contacts, knowledge documents and dummy data.

Only our own staff may call it (app/demo/admin_auth.py); with DEMO_ENABLED=false every
route answers 404. The logic lives in app/demo/records.py; this layer parses requests,
writes the audit trail, commits, and queues knowledge-base embedding after the commit.
"""
import io
import uuid
from datetime import datetime, timezone
from typing import Any

import yaml
from fastapi import (APIRouter, BackgroundTasks, Body, Depends, File, Form, HTTPException, Query,
                     UploadFile, status)
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, EmailStr, Field
from sqlalchemy import func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import get_settings
from app.db.base import get_db
from app.db.models import (DemoBatch, DemoContact, DemoImport, DemoIndustry, DemoProspect,
                           DemoProspectField, DemoSampleDataset)
from app.demo import records as svc
from app.demo.admin_auth import require_demo_admin
from app.demo.contacts import normalize_linkedin, tidy_name
from app.demo.extract import ExtractionError, extract_text
from app.demo.industries.definition import DefinitionError, compile_definition
from app.security.deps import Principal
from app.services.activity import append_activity_event

router = APIRouter(prefix="/admin/demo", tags=["demo-admin"])


def _fail(e: svc.DemoError):
    raise HTTPException(e.status, {"message": e.message, "errors": e.errors})


async def _read_upload(file: UploadFile, *, max_mb: int | None = None) -> bytes:
    limit = (max_mb or get_settings().demo_upload_max_file_mb) * 1024 * 1024
    data = await file.read(limit + 1)
    if len(data) > limit:
        raise HTTPException(status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                            f"file is larger than {limit // (1024 * 1024)} MB")
    return data


def _parse_definition_text(text: str) -> Any:
    try:
        return yaml.safe_load(text)
    except yaml.YAMLError as e:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY,
                            {"message": "not valid YAML/JSON", "errors": [str(e)[:500]]})


# ── industries ───────────────────────────────────────────────────────────────────────────

@router.get("/industries")
async def list_industries(p: Principal = Depends(require_demo_admin), db: AsyncSession = Depends(get_db)):
    rows = (await db.execute(select(DemoIndustry).order_by(DemoIndustry.key))).scalars().all()
    return [await svc.industry_readiness(db, r) for r in rows]


@router.get("/industries/{key}")
async def get_industry(key: str, p: Principal = Depends(require_demo_admin),
                       db: AsyncSession = Depends(get_db)):
    industry = await db.get(DemoIndustry, key)
    if industry is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "industry not found")
    try:
        definition = await svc.load_definition(db, key)
    except svc.DemoError as e:
        _fail(e)
    return {**await svc.industry_readiness(db, industry), "definition": definition.raw}


@router.post("/industries/validate")
async def validate_industry(file: UploadFile | None = File(None), text: str | None = Form(None),
                            p: Principal = Depends(require_demo_admin)):
    """Check a definition without saving it; returns every problem found."""
    raw = _parse_definition_text(text if text is not None else
                                 (await _read_upload(file, max_mb=1)).decode("utf-8", "replace")
                                 if file else "")
    try:
        d = compile_definition(raw if isinstance(raw, dict) else {})
    except DefinitionError as e:
        return {"valid": False, "errors": e.errors}
    return {"valid": True, "errors": [], "key": d.key, "facts": [f["key"] for f in d.facts_fields]}


@router.post("/industries", status_code=status.HTTP_201_CREATED)
async def publish_industry(file: UploadFile | None = File(None), text: str | None = Form(None),
                           activate: bool = Form(True),
                           p: Principal = Depends(require_demo_admin),
                           db: AsyncSession = Depends(get_db)):
    """Save a definition as the industry's next version (creates the industry if new).
    activate=true makes it the version new records use and publishes its dummy-data template."""
    if file is None and text is None:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "send `file` or `text`")
    raw = _parse_definition_text(text if text is not None else
                                 (await _read_upload(file, max_mb=1)).decode("utf-8", "replace"))
    try:
        definition, created = await svc.publish_industry_version(db, raw, actor=p.actor,
                                                                 activate=activate)
    except svc.DemoError as e:
        _fail(e)
    append_activity_event(db, p, "demo.industry_publish", resource_type="demo_industry",
                          resource_id=definition.key,
                          after={"version": definition.version, "activated": activate,
                                 "created": created})
    await db.commit()
    return {"key": definition.key, "version": definition.version, "activated": activate,
            "created": created}


@router.get("/industries/{key}/facts-template.xlsx")
async def facts_template(key: str, p: Principal = Depends(require_demo_admin),
                         db: AsyncSession = Depends(get_db)):
    """An empty workbook with the industry's tab and columns, for the research team."""
    import openpyxl
    try:
        definition = await svc.load_definition(db, key)
    except svc.DemoError as e:
        _fail(e)
    sheet = definition.sheet or {}
    headers = (list(dict.fromkeys(list(sheet.get("company", {}).values())
                                  + list((sheet.get("contacts") or {}).values())))
               + [f["sheet_column"] for f in definition.facts_fields
                  if f.get("sheet_column") and f["sheet_column"] not in
                  set(sheet.get("company", {}).values())])
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = (sheet.get("tab_names") or [definition.label])[0][:31]
    ws.append(headers)
    guide = wb.create_sheet("Field guide")
    guide.append(["Column", "Fact", "Required to go live"])
    for f in definition.facts_fields:
        guide.append([f.get("sheet_column", ""), f["label"], "Yes" if f.get("required") else ""])
    for rule in definition.minimum:
        guide.append(["", f"{rule['field']}: {rule.get('message', '')}", "Yes"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{key}-template.xlsx"'})


@router.get("/industries/{key}/dummy-data")
async def get_template_dataset(key: str, p: Principal = Depends(require_demo_admin),
                               db: AsyncSession = Depends(get_db)):
    row = (await db.execute(select(DemoSampleDataset).where(
        DemoSampleDataset.industry_key == key, DemoSampleDataset.prospect_id.is_(None),
        DemoSampleDataset.is_active.is_(True)))).scalar_one_or_none()
    if row is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "no dummy data template for this industry")
    return {"version": row.version, "industry_version": row.industry_version, "data": row.data}


@router.put("/industries/{key}/dummy-data")
async def put_template_dataset(key: str, data: dict = Body(...), dry_run: bool = Query(False),
                               p: Principal = Depends(require_demo_admin),
                               db: AsyncSession = Depends(get_db)):
    try:
        errors, row = await svc.save_template_dataset(db, key, data, dry_run=dry_run)
    except svc.DemoError as e:
        _fail(e)
    if errors:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY,
                            {"message": "the dummy data is invalid", "errors": errors})
    if dry_run:
        return {"valid": True, "saved": False}
    append_activity_event(db, p, "demo.template_data_save", resource_type="demo_industry",
                          resource_id=key, after={"version": row.version})
    await db.commit()
    return {"valid": True, "saved": True, "version": row.version}


@router.post("/industries/{key}/dummy-data/rollback/{version}")
async def rollback_template_dataset(key: str, version: int, p: Principal = Depends(require_demo_admin),
                                    db: AsyncSession = Depends(get_db)):
    try:
        row = await svc.rollback_dataset(db, industry_key=key, prospect_id=None, version=version)
    except svc.DemoError as e:
        _fail(e)
    append_activity_event(db, p, "demo.template_data_rollback", resource_type="demo_industry",
                          resource_id=key, after={"version": row.version})
    await db.commit()
    return {"active_version": row.version}


# ── batches ──────────────────────────────────────────────────────────────────────────────

class BatchIn(BaseModel):
    name: str = Field(min_length=1, max_length=120)


@router.post("/batches", status_code=status.HTTP_201_CREATED)
async def create_batch(body: BatchIn, p: Principal = Depends(require_demo_admin),
                       db: AsyncSession = Depends(get_db)):
    batch = DemoBatch(name=body.name.strip(), created_by=p.user_id)
    db.add(batch)
    await db.flush()
    append_activity_event(db, p, "demo.batch_create", resource_type="demo_batch",
                          resource_id=str(batch.id), after={"name": batch.name})
    await db.commit()
    return {"id": str(batch.id), "name": batch.name}


@router.get("/batches")
async def list_batches(p: Principal = Depends(require_demo_admin), db: AsyncSession = Depends(get_db)):
    rows = (await db.execute(select(DemoBatch, func.count(DemoContact.id))
                             .outerjoin(DemoContact, DemoContact.batch_id == DemoBatch.id)
                             .group_by(DemoBatch.id).order_by(DemoBatch.created_at.desc()))).all()
    return [{"id": str(b.id), "name": b.name, "contacts": int(n), "created_at": b.created_at}
            for b, n in rows]


# ── imports ──────────────────────────────────────────────────────────────────────────────

def _import_out(row: DemoImport, *, full: bool) -> dict:
    out = {"id": str(row.id), "industry_key": row.industry_key,
           "industry_version": row.industry_version, "source_name": row.source_name,
           "file_sha256": row.file_sha256, "status": row.status,
           "summary": (row.report or {}).get("summary"), "created_by": row.created_by,
           "created_at": row.created_at, "applied_by": row.applied_by, "applied_at": row.applied_at,
           "apply_result": row.apply_result}
    if full:
        out["report"] = row.report
    return out


@router.post("/imports/inspect")
async def inspect_import(file: UploadFile = File(...), p: Principal = Depends(require_demo_admin),
                         db: AsyncSession = Depends(get_db)):
    """What's in a workbook before importing: every tab, the industry it belongs to (or
    "new" for a prospect tab no industry claims yet), its header row and row count.
    Saves nothing."""
    from app.demo.importer import WorkbookError
    from app.demo.workbook import inspect_workbook
    data = await _read_upload(file)
    try:
        return inspect_workbook(data, await svc.all_current_definitions(db),
                                source_name=file.filename or "upload.xlsx")
    except WorkbookError as e:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, {"message": str(e), "errors": []})


@router.post("/industries/draft")
async def draft_industry(file: UploadFile = File(...), tab: str = Form(..., min_length=1, max_length=120),
                         p: Principal = Depends(require_demo_admin),
                         db: AsyncSession = Depends(get_db)):
    """A data-only industry definition drafted from a workbook tab no industry claims yet
    (its columns, and the workbook's Field guide when there is one). Saves nothing: the
    admin reviews the YAML and publishes it with POST /industries."""
    from app.demo.importer import WorkbookError
    from app.demo.settings import csv_list
    from app.demo.workbook import definition_yaml, draft_definition
    data = await _read_upload(file)
    try:
        raw, notes = draft_definition(data, tab, await svc.all_current_definitions(db),
                                      ignore_columns=csv_list(get_settings().demo_import_ignore_columns),
                                      source_name=file.filename or "upload.xlsx")
    except WorkbookError as e:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, {"message": str(e), "errors": []})
    try:
        compile_definition(raw)
        errors: list[str] = []
    except DefinitionError as e:
        errors = e.errors
    return {"key": raw["key"], "label": raw["label"], "definition": raw,
            "yaml": definition_yaml(raw), "notes": notes, "valid": not errors, "errors": errors}


def _skip_columns(value: str | None) -> list[str] | None:
    """`skip_columns` form field: a JSON list of column headers (headers may hold commas)."""
    if not value:
        return None
    import json
    try:
        parsed = json.loads(value)
    except ValueError:
        parsed = None
    if not (isinstance(parsed, list) and all(isinstance(x, str) for x in parsed)):
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY,
                            "skip_columns must be a JSON list of column headers")
    return parsed[:500]


@router.post("/imports", status_code=status.HTTP_201_CREATED)
async def preview_import(file: UploadFile = File(...), industry_key: str = Form(...),
                         tab: str | None = Form(None, max_length=120),
                         skip_columns: str | None = Form(None),
                         p: Principal = Depends(require_demo_admin),
                         db: AsyncSession = Depends(get_db)):
    """Dry run: reads the workbook, stores the report, writes nothing else. Apply it with
    POST /imports/{id}/apply.

    tab: the workbook tab to read (default: the industry's own tab names).
    skip_columns: JSON list of new columns to leave out. Every other column the industry
    doesn't map yet comes back in report.new_columns, and applying the import adds them to
    the industry definition (its next version)."""
    data = await _read_upload(file)
    try:
        row = await svc.preview_import(db, data=data, filename=file.filename or "upload.xlsx",
                                       industry_key=industry_key, actor=p.actor, tab=tab or None,
                                       skip_columns=_skip_columns(skip_columns))
    except svc.DemoError as e:
        _fail(e)
    append_activity_event(db, p, "demo.import_preview", resource_type="demo_import",
                          resource_id=str(row.id),
                          after={"file": row.source_name, "summary": row.report.get("summary")})
    await db.commit()
    return _import_out(row, full=True)


@router.get("/imports")
async def list_imports(limit: int = Query(20, ge=1, le=100), p: Principal = Depends(require_demo_admin),
                       db: AsyncSession = Depends(get_db)):
    rows = (await db.execute(select(DemoImport).order_by(DemoImport.created_at.desc()).limit(limit))).scalars()
    return [_import_out(r, full=False) for r in rows]


@router.get("/imports/{import_id}")
async def get_import(import_id: uuid.UUID, p: Principal = Depends(require_demo_admin),
                     db: AsyncSession = Depends(get_db)):
    row = await db.get(DemoImport, import_id)
    if row is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "import not found")
    return _import_out(row, full=True)


class ApplyIn(BaseModel):
    batch_id: uuid.UUID | None = None
    # The sheet rows to import (the ones left ticked in the preview); omitted = every row.
    rows: list[int] | None = Field(None, max_length=10000)


@router.post("/imports/{import_id}/apply")
async def apply_import(import_id: uuid.UUID, background: BackgroundTasks,
                       body: ApplyIn = Body(default_factory=ApplyIn),
                       p: Principal = Depends(require_demo_admin),
                       db: AsyncSession = Depends(get_db)):
    if body.batch_id is not None and await db.get(DemoBatch, body.batch_id) is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "batch not found")
    try:
        result, jobs = await svc.apply_import(db, import_id, actor=p.actor, batch_id=body.batch_id,
                                              rows=body.rows)
    except svc.DemoError as e:
        _fail(e)
    append_activity_event(db, p, "demo.import_apply", resource_type="demo_import",
                          resource_id=str(import_id),
                          after={k: v for k, v in result.items() if k != "records"})
    if result["columns_added"]:
        append_activity_event(db, p, "demo.industry_columns_add", resource_type="demo_industry",
                              resource_id=result.get("industry_key"),
                              after={"version": result["industry_version"],
                                     "columns": result["columns_added"],
                                     "import_id": str(import_id)})
    await db.commit()
    background.add_task(svc.run_ingest_jobs, jobs)
    return result


@router.post("/imports/{import_id}/discard")
async def discard_import(import_id: uuid.UUID, p: Principal = Depends(require_demo_admin),
                         db: AsyncSession = Depends(get_db)):
    row = await db.get(DemoImport, import_id)
    if row is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "import not found")
    if row.status != "previewed":
        raise HTTPException(status.HTTP_409_CONFLICT, f"this import is already {row.status}")
    row.status = "discarded"
    append_activity_event(db, p, "demo.import_discard", resource_type="demo_import",
                          resource_id=str(import_id))
    await db.commit()
    return {"status": "discarded"}


# ── records ──────────────────────────────────────────────────────────────────────────────

@router.get("/records")
async def list_records(industry_key: str | None = None, q: str | None = None,
                       minimum_met: bool | None = None,
                       limit: int = Query(50, ge=1, le=200), offset: int = Query(0, ge=0),
                       p: Principal = Depends(require_demo_admin), db: AsyncSession = Depends(get_db)):
    query = select(DemoProspect)
    if industry_key:
        query = query.where(DemoProspect.industry_key == industry_key)
    if q:
        like = f"%{q.strip()}%"
        sub = select(DemoProspectField.prospect_id).where(DemoProspectField.value.ilike(like))
        query = query.where(or_(DemoProspect.company_name.ilike(like), DemoProspect.city.ilike(like),
                                DemoProspect.id.in_(sub)))
    total = (await db.execute(select(func.count()).select_from(query.subquery()))).scalar_one()
    rows = (await db.execute(query.order_by(DemoProspect.company_name).limit(limit).offset(offset))).scalars()
    items = [await svc.record_summary(db, r, detail=False) for r in rows]
    if minimum_met is not None:     # computed from facts, so filtered after loading the page
        items = [i for i in items if i["minimum_met"] == minimum_met]
    return {"total": total, "limit": limit, "offset": offset, "items": items}


async def _record(db: AsyncSession, record_id: uuid.UUID) -> DemoProspect:
    try:
        return await svc.get_record(db, record_id)
    except svc.DemoError as e:
        _fail(e)


@router.get("/records/{record_id}")
async def get_record(record_id: uuid.UUID, p: Principal = Depends(require_demo_admin),
                     db: AsyncSession = Depends(get_db)):
    return await svc.record_summary(db, await _record(db, record_id), detail=True)


class RecordPatch(BaseModel):
    company_name: str | None = Field(None, min_length=1, max_length=200)
    company_name_pronunciation: str | None = Field(None, max_length=200)
    trigger_note: str | None = None
    timezone: str | None = None


@router.patch("/records/{record_id}")
async def patch_record(record_id: uuid.UUID, body: RecordPatch, background: BackgroundTasks,
                       p: Principal = Depends(require_demo_admin), db: AsyncSession = Depends(get_db)):
    from zoneinfo import ZoneInfo, ZoneInfoNotFoundError
    prospect = await _record(db, record_id)
    changes = body.model_dump(exclude_unset=True)
    if "timezone" in changes:
        try:
            ZoneInfo(changes["timezone"])
        except (ZoneInfoNotFoundError, ValueError, TypeError):
            raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "unknown timezone")
    before = {k: getattr(prospect, k) for k in changes}
    for k, v in changes.items():
        setattr(prospect, k, v)
    job = None
    if "company_name" in changes:
        definition = await svc.load_definition(db, prospect.industry_key, prospect.industry_version)
        job = await svc.refresh_company_facts_document(db, prospect, definition)
    await svc.provision_record(db, prospect)
    append_activity_event(db, p, "demo.record_update", resource_type="demo_record",
                          resource_id=str(record_id), before=before, after=changes)
    await db.commit()
    if job:
        background.add_task(svc.run_ingest_jobs, [job])
    return await svc.record_summary(db, prospect, detail=True)


class FactIn(BaseModel):
    value: str | None = None
    source_url: str | None = None


class FactsPatch(BaseModel):
    facts: dict[str, FactIn]


@router.patch("/records/{record_id}/facts")
async def patch_facts(record_id: uuid.UUID, body: FactsPatch, background: BackgroundTasks,
                      p: Principal = Depends(require_demo_admin), db: AsyncSession = Depends(get_db)):
    """Edit some facts (others are left alone). A null value clears a fact, which the agent
    then answers with "the sales team will confirm". Re-embeds the company facts document."""
    prospect = await _record(db, record_id)
    definition = await svc.load_definition(db, prospect.industry_key, prospect.industry_version)
    now = datetime.now(timezone.utc).isoformat()
    payload = {k: {"value": f.value, "source_url": f.source_url, "fetched_at": now}
               for k, f in body.facts.items()}
    try:
        changed = await svc.upsert_facts(db, prospect, definition, payload, replace_missing=False)
    except svc.DemoError as e:
        _fail(e)
    job = await svc.refresh_company_facts_document(db, prospect, definition)
    await svc.provision_record(db, prospect)
    append_activity_event(db, p, "demo.facts_update", resource_type="demo_record",
                          resource_id=str(record_id), after={"changed": changed})
    await db.commit()
    if job:
        background.add_task(svc.run_ingest_jobs, [job])
    return await svc.record_summary(db, prospect, detail=True)


@router.delete("/records/{record_id}")
async def delete_record(record_id: uuid.UUID, p: Principal = Depends(require_demo_admin),
                        db: AsyncSession = Depends(get_db)):
    """Erase a company record completely: its people, link, visitors, calls, leads and its
    whole demo tenant (agent, knowledge base, conversations). Cannot be undone."""
    from app.demo.erase import erase_records
    prospect = await _record(db, record_id)
    name = prospect.company_name
    try:
        counts = await erase_records(db, [prospect])
    except svc.DemoError as e:
        _fail(e)
    append_activity_event(db, p, "demo.record_erase", resource_type="demo_prospect",
                          resource_id=str(record_id), before={"company_name": name}, after=counts)
    await db.commit()
    return {"erased": True, **counts}


@router.post("/records/{record_id}/reindex")
async def reindex_record(record_id: uuid.UUID, background: BackgroundTasks,
                         p: Principal = Depends(require_demo_admin), db: AsyncSession = Depends(get_db)):
    """Re-embed the company facts document (e.g. after an embedding outage left it
    'processing' or 'failed')."""
    prospect = await _record(db, record_id)
    definition = await svc.load_definition(db, prospect.industry_key, prospect.industry_version)
    from app.db.models import KnowledgeDocument
    doc = (await db.execute(select(KnowledgeDocument).where(
        KnowledgeDocument.tenant_id == prospect.tenant_id,
        KnowledgeDocument.meta["demo_role"].astext == "company_facts"))).scalar_one_or_none()
    if doc is not None:
        doc.status = "processing"     # not "ready": refresh_company_facts_document rebuilds it
    job = await svc.refresh_company_facts_document(db, prospect, definition)
    await db.commit()
    if job:
        background.add_task(svc.run_ingest_jobs, [job])
    return {"queued": bool(job)}


# ── contacts ─────────────────────────────────────────────────────────────────────────────

class ContactIn(BaseModel):
    full_name: str = Field(min_length=1, max_length=160)
    role: str | None = Field(None, max_length=120)
    email: EmailStr | None = None
    linkedin_url: str | None = None
    batch_id: uuid.UUID | None = None


class ContactPatch(BaseModel):
    full_name: str | None = Field(None, min_length=1, max_length=160)
    role: str | None = Field(None, max_length=120)
    email: EmailStr | None = None
    linkedin_url: str | None = None
    batch_id: uuid.UUID | None = None
    status: str | None = Field(None, pattern="^(active|excluded)$")
    reviewed: bool | None = None


def _clean_linkedin(url: str | None) -> str | None:
    if not url:
        return None
    norm = normalize_linkedin(url)
    if norm is None:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "not a linkedin.com/in/ profile URL")
    return norm


@router.get("/records/{record_id}/contacts")
async def list_contacts(record_id: uuid.UUID, p: Principal = Depends(require_demo_admin),
                        db: AsyncSession = Depends(get_db)):
    await _record(db, record_id)
    rows = (await db.execute(select(DemoContact).where(DemoContact.prospect_id == record_id)
                             .order_by(DemoContact.full_name))).scalars()
    return [svc.contact_out(c) for c in rows]


@router.post("/records/{record_id}/contacts", status_code=status.HTTP_201_CREATED)
async def add_contact(record_id: uuid.UUID, body: ContactIn, p: Principal = Depends(require_demo_admin),
                      db: AsyncSession = Depends(get_db)):
    await _record(db, record_id)
    name = tidy_name(body.full_name.strip())
    email = str(body.email).lower() if body.email else None
    try:
        await svc.ensure_contact_unique(db, record_id, full_name=name, email=email)
    except svc.DemoError as e:
        _fail(e)
    c = DemoContact(prospect_id=record_id, batch_id=body.batch_id, full_name=name,
                    first_name=name.split()[0], role=body.role, email=email,
                    linkedin_url=_clean_linkedin(body.linkedin_url),
                    needs_review=not email, review_notes=None if email else "no email")
    db.add(c)
    await db.flush()
    append_activity_event(db, p, "demo.contact_add", resource_type="demo_contact",
                          resource_id=str(c.id), after={"record_id": str(record_id), "name": name})
    await db.commit()
    return svc.contact_out(c)


@router.patch("/contacts/{contact_id}")
async def patch_contact(contact_id: uuid.UUID, body: ContactPatch,
                        p: Principal = Depends(require_demo_admin), db: AsyncSession = Depends(get_db)):
    try:
        c = await svc.get_contact(db, contact_id)
    except svc.DemoError as e:
        _fail(e)
    changes = body.model_dump(exclude_unset=True)
    before = svc.contact_out(c)
    if "full_name" in changes:
        changes["full_name"] = tidy_name(changes["full_name"].strip())
    if changes.get("email"):
        changes["email"] = str(changes["email"]).lower()
    if "linkedin_url" in changes:
        changes["linkedin_url"] = _clean_linkedin(changes["linkedin_url"])
    if "batch_id" in changes and changes["batch_id"] and await db.get(DemoBatch, changes["batch_id"]) is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "batch not found")
    try:
        await svc.ensure_contact_unique(db, c.prospect_id, full_name=changes.get("full_name"),
                                        email=changes.get("email"), exclude=c.id)
    except svc.DemoError as e:
        _fail(e)
    reviewed = changes.pop("reviewed", None)
    for k, v in changes.items():
        setattr(c, k, v)
    if "full_name" in changes:
        c.first_name = c.full_name.split()[0]
    if reviewed is True:
        if not c.email:
            raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY,
                                "add an email before marking this contact reviewed")
        c.needs_review = False
        c.review_notes = f"reviewed by {p.label or p.actor}"
    elif reviewed is False:
        c.needs_review = True
    append_activity_event(db, p, "demo.contact_update", resource_type="demo_contact",
                          resource_id=str(c.id), before=before,
                          after={**{k: str(v) if isinstance(v, uuid.UUID) else v
                                    for k, v in changes.items()}, "reviewed": reviewed})
    try:
        await db.commit()
    except Exception as e:  # noqa: BLE001 — a race on the unique indexes
        await db.rollback()
        raise HTTPException(status.HTTP_409_CONFLICT, "that name or email is already taken") from e
    return svc.contact_out(c)


# ── knowledge documents ──────────────────────────────────────────────────────────────────

@router.get("/records/{record_id}/documents")
async def list_documents(record_id: uuid.UUID, p: Principal = Depends(require_demo_admin),
                         db: AsyncSession = Depends(get_db)):
    return await svc.list_kb_documents(db, await _record(db, record_id))


@router.post("/records/{record_id}/documents", status_code=status.HTTP_201_CREATED)
async def add_document(record_id: uuid.UUID, background: BackgroundTasks,
                       file: UploadFile | None = File(None), text: str | None = Form(None),
                       title: str | None = Form(None), doc_type: str = Form("brochure"),
                       source_url: str | None = Form(None),
                       p: Principal = Depends(require_demo_admin), db: AsyncSession = Depends(get_db)):
    """Add a document to the record's knowledge base: a file (PDF, HTML, TXT, Markdown) or
    pasted text. Embedded in the background; poll the documents list for its status."""
    prospect = await _record(db, record_id)
    if file is not None:
        try:
            content, kind = extract_text(await _read_upload(file), file.filename or "")
        except ExtractionError as e:
            raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, str(e))
        title = title or file.filename or "untitled"
    elif text and text.strip():
        content, kind = text.strip(), "text"
        title = title or "Pasted text"
    else:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "send a `file` or `text`")
    try:
        job = await svc.add_kb_document(db, prospect, title=title, text=content, doc_type=doc_type,
                                        kind=kind, source_uri=source_url)
    except svc.DemoError as e:
        _fail(e)
    append_activity_event(db, p, "demo.document_add", resource_type="demo_record",
                          resource_id=str(record_id), after={"title": title, "chars": len(content)})
    await db.commit()
    background.add_task(svc.run_ingest_jobs, [job])
    return {"id": str(job.document_id), "title": title, "status": "processing", "chars": len(content)}


@router.delete("/records/{record_id}/documents/{document_id}")
async def delete_document(record_id: uuid.UUID, document_id: uuid.UUID,
                          p: Principal = Depends(require_demo_admin), db: AsyncSession = Depends(get_db)):
    prospect = await _record(db, record_id)
    try:
        title = await svc.delete_kb_document(db, prospect, document_id)
    except svc.DemoError as e:
        _fail(e)
    append_activity_event(db, p, "demo.document_delete", resource_type="demo_record",
                          resource_id=str(record_id), before={"title": title})
    await db.commit()
    from app.engine.rag.retrieval import invalidate_tenant_knowledge_cache
    await invalidate_tenant_knowledge_cache(prospect.tenant_id)
    return {"deleted": True}


# ── dummy data per record ────────────────────────────────────────────────────────────────

@router.get("/records/{record_id}/dummy-data")
async def get_record_dataset(record_id: uuid.UUID, p: Principal = Depends(require_demo_admin),
                             db: AsyncSession = Depends(get_db)):
    prospect = await _record(db, record_id)
    definition = await svc.load_definition(db, prospect.industry_key, prospect.industry_version)
    _, data, source, version = await svc.effective_dataset(db, prospect, definition)
    errors = svc.check_dataset(definition, data, await svc.load_facts(db, prospect, definition)) \
        if data is not None else ["no dummy data"]
    return {"source": source, "version": version, "data": data, "errors": errors}


@router.put("/records/{record_id}/dummy-data")
async def put_record_dataset(record_id: uuid.UUID, data: dict = Body(...), dry_run: bool = Query(False),
                             p: Principal = Depends(require_demo_admin), db: AsyncSession = Depends(get_db)):
    """Give this record its own dummy data instead of the industry template."""
    prospect = await _record(db, record_id)
    try:
        errors, row = await svc.save_record_dataset(db, prospect, data, dry_run=dry_run)
    except svc.DemoError as e:
        _fail(e)
    if errors:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY,
                            {"message": "the dummy data is invalid for this record", "errors": errors})
    if dry_run:
        return {"valid": True, "saved": False}
    await svc.provision_record(db, prospect)
    append_activity_event(db, p, "demo.record_data_save", resource_type="demo_record",
                          resource_id=str(record_id), after={"version": row.version})
    await db.commit()
    return {"valid": True, "saved": True, "version": row.version}


@router.delete("/records/{record_id}/dummy-data")
async def clear_record_dataset(record_id: uuid.UUID, p: Principal = Depends(require_demo_admin),
                               db: AsyncSession = Depends(get_db)):
    """Go back to the industry template (the record's versions are kept for rollback)."""
    prospect = await _record(db, record_id)
    cleared = await svc.clear_record_dataset(db, prospect)
    await svc.provision_record(db, prospect)
    if cleared:
        append_activity_event(db, p, "demo.record_data_clear", resource_type="demo_record",
                              resource_id=str(record_id))
    await db.commit()
    return {"source": "template", "cleared": cleared}


@router.post("/records/{record_id}/dummy-data/rollback/{version}")
async def rollback_record_dataset(record_id: uuid.UUID, version: int,
                                  p: Principal = Depends(require_demo_admin),
                                  db: AsyncSession = Depends(get_db)):
    prospect = await _record(db, record_id)
    try:
        row = await svc.rollback_dataset(db, industry_key=prospect.industry_key,
                                         prospect_id=prospect.id, version=version)
    except svc.DemoError as e:
        _fail(e)
    await svc.provision_record(db, prospect)
    append_activity_event(db, p, "demo.record_data_rollback", resource_type="demo_record",
                          resource_id=str(record_id), after={"version": row.version})
    await db.commit()
    return {"active_version": row.version}


@router.get("/records/{record_id}/dummy-data/preview")
async def preview_record_dataset(record_id: uuid.UUID, at: datetime | None = None,
                                 p: Principal = Depends(require_demo_admin),
                                 db: AsyncSession = Depends(get_db)):
    """The dummy data exactly as a call at `at` (default: now) would get it."""
    prospect = await _record(db, record_id)
    at = at or datetime.now(timezone.utc)
    if at.tzinfo is None:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "`at` needs a timezone offset")
    try:
        return await svc.preview_dataset(db, prospect, at)
    except svc.DemoError as e:
        _fail(e)


# ── demo agent: provisioning and a text test chat ────────────────────────────────────────

@router.post("/records/{record_id}/agent")
async def provision_agent(record_id: uuid.UUID, p: Principal = Depends(require_demo_admin),
                          db: AsyncSession = Depends(get_db)):
    """(Re)build the record's demo agent from its facts and dummy data. Normally automatic
    (import, edits, every call start); exposed for checking a record's readiness."""
    from app.demo.agent import provision_demo_agent
    prospect = await _record(db, record_id)
    try:
        out = await provision_demo_agent(db, prospect)
    except svc.DemoError as e:
        _fail(e)
    append_activity_event(db, p, "demo.agent_provision", resource_type="demo_record",
                          resource_id=str(record_id), after={"changed": out["changed"]})
    await db.commit()
    return out


@router.post("/records/{record_id}/chat", status_code=status.HTTP_201_CREATED)
async def start_chat(record_id: uuid.UUID, p: Principal = Depends(require_demo_admin),
                     db: AsyncSession = Depends(get_db)):
    """Start a text chat with the record's demo agent, as a fresh sample lead — the QA way
    to test a demo before any link goes out."""
    from app.demo.agent import start_demo_conversation
    prospect = await _record(db, record_id)
    try:
        conv, customer, greeting = await start_demo_conversation(db, prospect, channel="chat")
    except svc.DemoError as e:
        _fail(e)
    await db.commit()
    return {"session_id": str(conv.id), "opening_line": greeting,
            "caller": {"id": str(customer.id), "name": customer.full_name}}


class ChatTurn(BaseModel):
    text: str = Field(min_length=1, max_length=2000)


async def _demo_conversation(db: AsyncSession, prospect: DemoProspect, session_id: uuid.UUID):
    from app.db.models import Conversation
    conv = (await db.execute(select(Conversation).where(
        Conversation.id == session_id, Conversation.tenant_id == prospect.tenant_id))).scalar_one_or_none()
    if conv is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "chat not found for this record")
    return conv


async def _transcript_messages(db: AsyncSession, conversation_id: uuid.UUID) -> list[dict]:
    """What the caller said and what the agent said back, in order. The engine stores its
    internals in the same conversation — tool results (role "tool") and planner actions
    (role "system", e.g. "action:respond") — and those are not part of a transcript."""
    from app.db.models import Message
    msgs = (await db.execute(select(Message).where(
        Message.conversation_id == conversation_id, Message.role.in_(("customer", "agent")))
        .order_by(Message.created_at))).scalars().all()
    return [{"role": m.role, "content": m.content, "at": m.created_at} for m in msgs]


@router.post("/records/{record_id}/chat/{session_id}/turns")
async def chat_turn(record_id: uuid.UUID, session_id: uuid.UUID, body: ChatTurn,
                    p: Principal = Depends(require_demo_admin), db: AsyncSession = Depends(get_db)):
    """One turn through the real engine, in the record's demo tenant. Admins see the debug
    envelope (tools, rules); prospects never do."""
    from app.engine.planner.turn_runner import process_turn
    prospect = await _record(db, record_id)
    conv = await _demo_conversation(db, prospect, session_id)
    if conv.status != "active":
        raise HTTPException(status.HTTP_409_CONFLICT, f"chat is {conv.status}")
    result = await process_turn(db, tenant_id=prospect.tenant_id, conversation=conv,
                                user_input=body.text, actor=p.actor)
    return {"response_text": result.response_text, "executed_tools": result.executed_tools,
            "state_changed_to": result.state_changed_to, "end_call": result.end_call,
            "debug": result.debug}


@router.get("/records/{record_id}/chat/{session_id}")
async def chat_transcript(record_id: uuid.UUID, session_id: uuid.UUID,
                          p: Principal = Depends(require_demo_admin), db: AsyncSession = Depends(get_db)):
    from app.db.models import Appointment
    prospect = await _record(db, record_id)
    conv = await _demo_conversation(db, prospect, session_id)
    appts = (await db.execute(select(Appointment).where(
        Appointment.tenant_id == prospect.tenant_id,
        Appointment.customer_id == conv.customer_id))).scalars().all()
    return {"session_id": str(conv.id), "status": conv.status,
            "messages": await _transcript_messages(db, conv.id),
            "bookings": [{"starts_at": a.starts_at, "status": a.status} for a in appts]}


# ── company links ────────────────────────────────────────────────────────────────────────
#
# One private link per company record; each person it is sent to gets their own recipient
# code in their copy of the URL (app/demo/links.py), so opens, calls and questions are
# attributed to a person. URLs can be shown again (they are rebuilt from the encrypted
# token); every time they are, the audit log records who looked.

class LinkIn(BaseModel):
    replace: bool = False            # revoke the record's current link and make a new one
    ttl_days: int | None = Field(None, ge=1, le=90)
    max_calls_per_day: int | None = Field(None, ge=1)
    max_call_seconds: int | None = Field(None, ge=30)
    max_total_minutes: int | None = Field(None, ge=1)
    max_concurrent_calls: int | None = Field(None, ge=1)


_LINK_LIMITS = {"max_calls_per_day", "max_call_seconds", "max_total_minutes", "max_concurrent_calls"}


async def _batches(db: AsyncSession, ids) -> dict:
    ids = {i for i in ids if i}
    if not ids:
        return {}
    return {b.id: b for b in (await db.execute(select(DemoBatch).where(DemoBatch.id.in_(ids)))).scalars()}


async def _link_urls(db: AsyncSession, link, prospect: DemoProspect) -> dict:
    """Every URL of a link: the plain company link and each person's own copies."""
    from app.demo.links import campaign_for, link_url, recipient_urls, recipients_of, reveal_token
    raw = reveal_token(link)
    people = await recipients_of(db, link.id)
    batches = await _batches(db, [c.batch_id for _, c in people])
    return {
        "company_url": link_url(raw, campaign=campaign_for(prospect, None)),
        "recipients": [{"recipient_id": str(r.id), "ref_code": r.ref_code,
                        "last_exported_at": r.last_exported_at,
                        "contact": {"id": str(c.id), "name": c.full_name, "email": c.email,
                                    "role": c.role, "status": c.status},
                        "urls": recipient_urls(raw, r.ref_code, campaign_for(prospect, batches.get(c.batch_id)))}
                       for r, c in people],
    }


@router.post("/records/{record_id}/link", status_code=status.HTTP_201_CREATED)
async def create_record_link(record_id: uuid.UUID, body: LinkIn = Body(default_factory=LinkIn),
                             p: Principal = Depends(require_demo_admin), db: AsyncSession = Depends(get_db)):
    """The company link of a record, with a recipient code for each active person there.
    Returns the existing active link unless replace=true (which stops the old URLs)."""
    from app.demo.links import ensure_recipient, issue_link, link_out
    try:
        prospect = await svc.get_record(db, record_id)
        issued = await issue_link(db, prospect, replace=body.replace, ttl_days=body.ttl_days,
                                  limits=body.model_dump(include=_LINK_LIMITS), created_by=p.user_id)
        for c in (await db.execute(select(DemoContact).where(
                DemoContact.prospect_id == prospect.id, DemoContact.status == "active"))).scalars():
            await ensure_recipient(db, issued.link, c, created_by=p.user_id)
        urls = await _link_urls(db, issued.link, prospect)
    except svc.DemoError as e:
        _fail(e)
    if issued.created:
        append_activity_event(db, p, "demo.link_create", resource_type="demo_link",
                              resource_id=str(issued.link.id),
                              after={"record_id": str(prospect.id), "replace": body.replace,
                                     "expires_at": issued.link.expires_at.isoformat()})
    append_activity_event(db, p, "demo.link_reveal", resource_type="demo_link", resource_id=str(issued.link.id))
    await db.commit()
    return {**link_out(issued.link, prospect=prospect), "created": issued.created, **urls}


@router.get("/links/{link_id}/urls")
async def link_urls(link_id: uuid.UUID, p: Principal = Depends(require_demo_admin),
                    db: AsyncSession = Depends(get_db)):
    """The URLs of a link (the company link and each person's copy). Audited."""
    from app.demo.links import get_link, link_out
    try:
        link = await get_link(db, link_id)
        prospect = await db.get(DemoProspect, link.prospect_id)
        urls = await _link_urls(db, link, prospect)
    except svc.DemoError as e:
        _fail(e)
    append_activity_event(db, p, "demo.link_reveal", resource_type="demo_link", resource_id=str(link.id))
    await db.commit()
    return {**link_out(link, prospect=prospect), **urls}


class RecipientIn(BaseModel):
    contact_id: uuid.UUID


@router.post("/links/{link_id}/recipients", status_code=status.HTTP_201_CREATED)
async def add_recipient(link_id: uuid.UUID, body: RecipientIn, p: Principal = Depends(require_demo_admin),
                        db: AsyncSession = Depends(get_db)):
    """Send a company's link to one more person: anyone at that company added later, or a
    person at another company (e.g. a group's head office). Their own URLs come back."""
    from app.demo.links import campaign_for, effective_status, ensure_recipient, get_link, recipient_urls, reveal_token
    try:
        link = await get_link(db, link_id)
        if effective_status(link) == "revoked":
            raise svc.Conflict("this link is revoked; make a new one for the company")
        contact = await svc.get_contact(db, body.contact_id)
        prospect = await db.get(DemoProspect, link.prospect_id)
        r = await ensure_recipient(db, link, contact, created_by=p.user_id)
        batch = (await _batches(db, [contact.batch_id])).get(contact.batch_id)
        urls = recipient_urls(reveal_token(link), r.ref_code, campaign_for(prospect, batch))
    except svc.DemoError as e:
        _fail(e)
    append_activity_event(db, p, "demo.link_recipient_add", resource_type="demo_link", resource_id=str(link.id),
                          after={"contact_id": str(contact.id), "recipient_id": str(r.id)})
    await db.commit()
    return {"recipient_id": str(r.id), "ref_code": r.ref_code,
            "contact": {"id": str(contact.id), "name": contact.full_name, "email": contact.email},
            "urls": urls}


class BulkLinksIn(BaseModel):
    record_ids: list[uuid.UUID] | None = None      # none + no batch/industry = every record
    batch_id: uuid.UUID | None = None              # only records with people in this batch, and only those people
    industry_key: str | None = None
    regenerate: bool = False                       # new links (old URLs stop working)
    ttl_days: int | None = Field(None, ge=1, le=90)


def _csv_fields(steps: list[str]) -> list[str]:
    return ["first_name", "full_name", "email", "role", "company", "record_name", "industry", "link",
            *[f"link_{s}" for s in steps], "company_link", "expires_on", "needs_review", "review_notes",
            "other_company", "prompt_1", "prompt_2", "ref_code", "record_id", "link_id", "contact_id"]


@router.post("/links/generate")
async def generate_links(body: BulkLinksIn, format: str = Query("json", pattern="^(json|csv)$"),
                         p: Principal = Depends(require_demo_admin), db: AsyncSession = Depends(get_db)):
    """One link per company record and one row per person it goes to, for the email tool.
    A record keeps its current link (same URLs as last time) unless regenerate=true. Each
    person's URL carries their own recipient code, the email step and UTM tags; `link` is
    the first step's URL. Exporting marks those people as emailed."""
    import csv

    from app.demo.links import (campaign_for, ensure_recipient, export_steps, issue_link, link_url,
                                recipient_urls, recipients_of)
    q = select(DemoProspect)
    if body.record_ids:
        q = q.where(DemoProspect.id.in_(body.record_ids))
    if body.industry_key:
        q = q.where(DemoProspect.industry_key == body.industry_key)
    if body.batch_id:
        q = q.where(DemoProspect.id.in_(select(DemoContact.prospect_id).where(DemoContact.batch_id == body.batch_id)))
    prospects = (await db.execute(q.order_by(DemoProspect.company_name))).scalars().all()

    now = datetime.now(timezone.utc)
    steps = export_steps()
    rows, links, skipped = [], [], []
    for prospect in prospects:
        try:
            issued = await issue_link(db, prospect, replace=body.regenerate, ttl_days=body.ttl_days,
                                      created_by=p.user_id)
        except svc.DemoError as e:
            skipped.append({"record_id": str(prospect.id), "company": prospect.company_name, "reason": e.message})
            continue
        own = (await db.execute(select(DemoContact).where(DemoContact.prospect_id == prospect.id)
                                .order_by(DemoContact.full_name))).scalars().all()
        for c in own:
            if body.batch_id and c.batch_id != body.batch_id:
                continue
            if c.status != "active":
                skipped.append({"record_id": str(prospect.id), "company": prospect.company_name,
                                "contact_id": str(c.id), "name": c.full_name, "reason": "excluded"})
                continue
            await ensure_recipient(db, issued.link, c, created_by=p.user_id)
        definition = await svc.load_definition(db, prospect.industry_key, prospect.industry_version)
        prompts = list(definition.raw.get("prompts") or [])
        facts = await svc.load_facts(db, prospect, definition)
        rn = (definition.sheet or {}).get("record_name_field")
        record_name = facts[rn].value if rn and not facts[rn].is_blank else ""
        people = [(r, c) for r, c in await recipients_of(db, issued.link.id)
                  if c.status == "active" and (not body.batch_id or c.batch_id == body.batch_id)]
        batches = await _batches(db, [c.batch_id for _, c in people])
        company_link = link_url(issued.raw, campaign=campaign_for(prospect, None))
        links.append({"record_id": str(prospect.id), "company": prospect.company_name,
                      "link_id": str(issued.link.id), "created": issued.created, "people": len(people),
                      "expires_at": issued.link.expires_at})
        for r, c in people:
            urls = recipient_urls(issued.raw, r.ref_code, campaign_for(prospect, batches.get(c.batch_id)))
            r.last_exported_at = now
            rows.append({
                "first_name": c.first_name or "", "full_name": c.full_name, "email": c.email or "",
                "role": c.role or "", "company": prospect.company_name, "record_name": record_name,
                "industry": definition.label, "link": urls[steps[0]],
                **{f"link_{s}": u for s, u in urls.items()}, "company_link": company_link,
                "expires_on": issued.link.expires_at.date().isoformat(),
                "needs_review": "yes" if c.needs_review else "", "review_notes": c.review_notes or "",
                "other_company": "yes" if c.prospect_id != prospect.id else "",
                "prompt_1": prompts[0] if prompts else "", "prompt_2": prompts[1] if len(prompts) > 1 else "",
                "ref_code": r.ref_code, "record_id": str(prospect.id), "link_id": str(issued.link.id),
                "contact_id": str(c.id)})
    append_activity_event(db, p, "demo.links_generate", resource_type="demo_link",
                          after={"links": len(links), "new_links": sum(1 for x in links if x["created"]),
                                 "rows": len(rows), "skipped": len(skipped), "regenerate": body.regenerate})
    await db.commit()
    if format == "csv":
        buf = io.StringIO()
        writer = csv.DictWriter(buf, fieldnames=_csv_fields(steps))
        writer.writeheader()
        writer.writerows(rows)
        stamp = now.strftime("%Y%m%d-%H%M")
        return StreamingResponse(
            io.BytesIO(buf.getvalue().encode("utf-8-sig")), media_type="text/csv",
            headers={"Content-Disposition": f'attachment; filename="demo-links-{stamp}.csv"',
                     "Cache-Control": "no-store"})
    return {"links": links, "rows": rows, "skipped": skipped}


@router.get("/links")
async def list_links(record_id: uuid.UUID | None = None, industry_key: str | None = None,
                     include_inactive: bool = False, q: str | None = Query(None, max_length=100),
                     limit: int = Query(200, ge=1, le=500), p: Principal = Depends(require_demo_admin),
                     db: AsyncSession = Depends(get_db)):
    """Company links with who opened them (no tokens: only a 6-character prefix)."""
    from app.db.models import DemoLink
    from app.demo.activity import link_stats
    from app.demo.links import link_out
    query = select(DemoLink, DemoProspect).join(DemoProspect, DemoProspect.id == DemoLink.prospect_id)
    if record_id:
        query = query.where(DemoLink.prospect_id == record_id)
    if industry_key:
        query = query.where(DemoProspect.industry_key == industry_key)
    if not include_inactive and not record_id:
        query = query.where(DemoLink.status == "active")
    if q:
        query = query.where(DemoProspect.company_name.ilike(f"%{q}%"))
    rows = (await db.execute(query.order_by(DemoLink.created_at.desc()).limit(limit))).all()
    stats = await link_stats(db, [link.id for link, _ in rows])
    return [{**link_out(link, prospect=pr), "stats": stats[link.id]} for link, pr in rows]


@router.get("/links/{link_id}/activity")
async def link_activity(link_id: uuid.UUID, p: Principal = Depends(require_demo_admin),
                        db: AsyncSession = Depends(get_db)):
    """Everything that happened on one company link: who it was sent to and what each of
    them did, visitors we couldn't tie to a person, calls, the timeline, and every question."""
    from app.demo import activity
    from app.demo.links import get_link, link_out
    try:
        link = await get_link(db, link_id)
    except svc.DemoError as e:
        _fail(e)
    prospect = await db.get(DemoProspect, link.prospect_id)
    return {"link": link_out(link, prospect=prospect), **(await activity.link_activity(db, link))}


@router.get("/questions")
async def list_questions(record_id: uuid.UUID | None = None, link_id: uuid.UUID | None = None,
                         industry_key: str | None = None, q: str | None = Query(None, max_length=100),
                         limit: int = Query(100, ge=1, le=500), offset: int = Query(0, ge=0),
                         p: Principal = Depends(require_demo_admin), db: AsyncSession = Depends(get_db)):
    """What prospects said to the demo agents, newest first, with who said it and the
    agent's reply. `q` searches the text ("discount", "RERA")."""
    from app.demo.activity import questions
    return await questions(db, record_id=record_id, link_id=link_id, industry_key=industry_key,
                           contains=q, limit=limit, offset=offset)


class RevokeIn(BaseModel):
    reason: str | None = Field(None, max_length=300)


@router.post("/links/{link_id}/revoke")
async def revoke_link(link_id: uuid.UUID, body: RevokeIn = Body(default_factory=RevokeIn),
                      p: Principal = Depends(require_demo_admin), db: AsyncSession = Depends(get_db)):
    """Stop a company link for everyone it was sent to."""
    from app.demo.links import get_link, link_out, now_utc
    try:
        link = await get_link(db, link_id)
    except svc.DemoError as e:
        _fail(e)
    if link.status != "revoked":
        link.status, link.revoked_at = "revoked", now_utc()
        link.revoked_reason = body.reason or "revoked by an admin"
        append_activity_event(db, p, "demo.link_revoke", resource_type="demo_link",
                              resource_id=str(link.id), after={"reason": link.revoked_reason})
        await db.commit()
    return link_out(link)


class ExtendIn(BaseModel):
    days: int = Field(ge=1, le=90)


@router.post("/links/{link_id}/extend")
async def extend_link(link_id: uuid.UUID, body: ExtendIn, p: Principal = Depends(require_demo_admin),
                      db: AsyncSession = Depends(get_db)):
    """Push the expiry out by `days` from now (or from the current expiry, if later). A
    revoked link stays revoked: generate a new one instead."""
    from datetime import timedelta

    from app.demo.links import get_link, link_out, now_utc
    try:
        link = await get_link(db, link_id)
    except svc.DemoError as e:
        _fail(e)
    if link.status == "revoked":
        raise HTTPException(status.HTTP_409_CONFLICT, "a revoked link can't be extended; generate a new one")
    before = link.expires_at
    link.expires_at = max(link.expires_at, now_utc()) + timedelta(days=body.days)
    link.status = "active"
    append_activity_event(db, p, "demo.link_extend", resource_type="demo_link", resource_id=str(link.id),
                          before={"expires_at": before.isoformat()},
                          after={"expires_at": link.expires_at.isoformat()})
    await db.commit()
    return link_out(link)


# ── activity: demo calls, transcripts, funnel ────────────────────────────────────────────

@router.get("/sessions")
async def list_sessions(record_id: uuid.UUID | None = None, link_id: uuid.UUID | None = None,
                        limit: int = Query(50, ge=1, le=200), p: Principal = Depends(require_demo_admin),
                        db: AsyncSession = Depends(get_db)):
    """Demo calls, newest first, with who made them and how they ended."""
    from app.db.models import DemoLink, DemoSession
    from app.demo.activity import session_people
    q = (select(DemoSession, DemoLink, DemoProspect)
         .join(DemoLink, DemoLink.id == DemoSession.demo_link_id)
         .join(DemoProspect, DemoProspect.id == DemoLink.prospect_id))
    if record_id:
        q = q.where(DemoLink.prospect_id == record_id)
    if link_id:
        q = q.where(DemoSession.demo_link_id == link_id)
    rows = (await db.execute(q.order_by(DemoSession.created_at.desc()).limit(limit))).all()
    people = await session_people(db, [s for s, _, _ in rows])
    out = []
    for s, link, pr in rows:
        who = people.who(s.demo_visitor_id, s.recipient_id)
        out.append({"id": str(s.id), "link_id": str(link.id), "who": who,
                    "contact": who["name"], "email": who["email"],
                    "company": pr.company_name, "record_id": str(pr.id), "created_at": s.created_at,
                    "started_at": s.started_at, "ended_at": s.ended_at, "duration_seconds": s.duration_seconds,
                    "outcome": s.outcome, "end_reason": s.end_reason,
                    "has_transcript": s.conversation_id is not None})
    return out


@router.get("/sessions/{session_id}/transcript")
async def session_transcript(session_id: uuid.UUID, p: Principal = Depends(require_demo_admin),
                             db: AsyncSession = Depends(get_db)):
    from app.db.models import Appointment, Conversation, DemoSession
    from app.demo.activity import session_people
    s = await db.get(DemoSession, session_id)
    if s is None or s.conversation_id is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "no transcript for this session")
    conv = await db.get(Conversation, s.conversation_id)
    appts = []
    if conv is not None and conv.customer_id is not None:
        appts = (await db.execute(select(Appointment).where(
            Appointment.tenant_id == conv.tenant_id,
            Appointment.customer_id == conv.customer_id))).scalars().all()
    people = await session_people(db, [s])
    return {"session_id": str(s.id), "outcome": s.outcome, "duration_seconds": s.duration_seconds,
            "who": people.who(s.demo_visitor_id, s.recipient_id),
            "messages": await _transcript_messages(db, s.conversation_id),
            "bookings": [{"starts_at": a.starts_at, "status": a.status} for a in appts]}


@router.get("/metrics")
async def demo_metrics(industry_key: str | None = None, p: Principal = Depends(require_demo_admin),
                       db: AsyncSession = Depends(get_db)):
    """The funnel (docs §9.5), per company and per person: companies with a link → opened
    → talked → goal reached; people emailed → opened → talked; plus usage and spend."""
    from sqlalchemy import distinct

    from app.db.models import DemoLink, DemoLinkRecipient, DemoSession, DemoVisitor
    from app.demo.limits import spend_today_usd

    def scoped(q, link_col):
        if industry_key:
            q = q.join(DemoProspect, DemoProspect.id == link_col).where(DemoProspect.industry_key == industry_key)
        return q

    async def one(q):
        return (await db.execute(q)).scalar_one()

    now = datetime.now(timezone.utc)
    records_q = select(func.count(DemoProspect.id))
    if industry_key:
        records_q = records_q.where(DemoProspect.industry_key == industry_key)
    contacts_q = select(func.count(DemoContact.id)).where(DemoContact.status == "active")
    if industry_key:
        contacts_q = contacts_q.join(DemoProspect, DemoProspect.id == DemoContact.prospect_id).where(
            DemoProspect.industry_key == industry_key)
    started = DemoSession.started_at.is_not(None)
    viewed = DemoVisitor.first_viewed_at.is_not(None)
    via_link = lambda q, col: scoped(q.join(DemoLink, DemoLink.id == col), DemoLink.prospect_id)  # noqa: E731

    links_active = await one(scoped(select(func.count(DemoLink.id)).where(
        DemoLink.status == "active", DemoLink.expires_at > now), DemoLink.prospect_id))
    companies_opened = await one(via_link(select(func.count(distinct(DemoLink.prospect_id))).select_from(DemoVisitor)
                                          .where(viewed), DemoVisitor.demo_link_id))
    companies_talked = await one(via_link(select(func.count(distinct(DemoLink.prospect_id))).select_from(DemoSession)
                                          .where(started), DemoSession.demo_link_id))
    companies_2min = await one(via_link(select(func.count(distinct(DemoLink.prospect_id))).select_from(DemoSession)
                                        .where(DemoSession.duration_seconds >= 120), DemoSession.demo_link_id))
    people_emailed = await one(via_link(select(func.count(distinct(DemoLinkRecipient.contact_id)))
                                        .select_from(DemoLinkRecipient)
                                        .where(DemoLinkRecipient.last_exported_at.is_not(None)),
                                        DemoLinkRecipient.demo_link_id))
    people_opened = await one(via_link(select(func.count(distinct(DemoLinkRecipient.contact_id)))
                                       .select_from(DemoVisitor)
                                       .join(DemoLinkRecipient, DemoLinkRecipient.id == DemoVisitor.recipient_id)
                                       .where(viewed), DemoVisitor.demo_link_id))
    people_talked = await one(via_link(select(func.count(distinct(DemoLinkRecipient.contact_id)))
                                       .select_from(DemoSession)
                                       .join(DemoLinkRecipient, DemoLinkRecipient.id == DemoSession.recipient_id)
                                       .where(started), DemoSession.demo_link_id))
    unknown_visitors = await one(via_link(select(func.count(DemoVisitor.id)).where(
        viewed, DemoVisitor.recipient_id.is_(None)), DemoVisitor.demo_link_id))
    calls = (await db.execute(via_link(select(func.count(DemoSession.id),
                                              func.coalesce(func.sum(DemoSession.duration_seconds), 0))
                                       .where(started), DemoSession.demo_link_id))).one()
    goals = await one(via_link(select(func.count(DemoSession.id)).where(
        DemoSession.outcome == "goal_completed"), DemoSession.demo_link_id))
    return {"records": await one(records_q), "contacts": await one(contacts_q),
            "links_active": links_active,
            "companies_opened": companies_opened, "companies_talked": companies_talked,
            "companies_talked_2min": companies_2min,
            "people_emailed": people_emailed, "people_opened": people_opened, "people_talked": people_talked,
            "unknown_visitors": unknown_visitors,
            "calls": int(calls[0]), "call_minutes": round(int(calls[1]) / 60, 1), "visits_booked": goals,
            "spend_today_usd": float(await spend_today_usd(db, now)),
            # Company level, as the spec measures it (a prospect = a company).
            "open_rate": round(companies_opened / links_active, 3) if links_active else None,
            "talk_rate": round(companies_talked / companies_opened, 3) if companies_opened else None,
            "people_open_rate": round(people_opened / people_emailed, 3) if people_emailed else None}


# ── leads (Contact us / Schedule a call; the team works them in the Google Sheet) ───────

def _lead_out(lead, contact_name: str | None) -> dict:
    from app.demo.leads.mailer import email_on
    from app.demo.leads.service import KIND_LABEL, sheet_on
    s = get_settings()
    pending = ((sheet_on() and lead.sheet_synced_at is None) or (email_on() and lead.email_notified_at is None))
    return {"id": str(lead.id), "lead_id": lead.sheet_row_key, "kind": lead.kind,
            "kind_label": KIND_LABEL.get(lead.kind, lead.kind), "name": lead.name, "email": lead.email,
            "phone": lead.phone, "company": lead.company, "role": lead.role,
            "preferred_time": lead.preferred_time, "message": lead.message,
            "record_id": str(lead.prospect_id) if lead.prospect_id else None,
            "link_id": str(lead.demo_link_id) if lead.demo_link_id else None,
            "on_our_list": contact_name, "last_session_id": str(lead.last_session_id) if lead.last_session_id else None,
            "call_status": lead.call_status, "called_by": lead.called_by, "notes": lead.notes,
            "created_at": lead.created_at, "sheet_synced_at": lead.sheet_synced_at,
            "email_notified_at": lead.email_notified_at, "notify_attempts": lead.notify_attempts,
            "last_notify_error": lead.last_notify_error,
            "delivery": "failed" if pending and lead.notify_attempts >= s.demo_notify_max_attempts
            else "pending" if pending else "delivered"}


@router.get("/leads")
async def list_leads(q: str | None = Query(None, max_length=100), kind: str | None = None,
                     limit: int = Query(200, ge=1, le=1000), p: Principal = Depends(require_demo_admin),
                     db: AsyncSession = Depends(get_db)):
    """Leads from the demo pages, newest first. Status / Called by / Notes are the team's,
    copied back from the Google Sheet."""
    from app.db.models import DemoLead
    query = select(DemoLead, DemoContact.full_name).outerjoin(DemoContact, DemoContact.id == DemoLead.contact_id)
    if kind:
        query = query.where(DemoLead.kind == kind)
    if q:
        like = f"%{q}%"
        query = query.where(or_(DemoLead.name.ilike(like), DemoLead.email.ilike(like),
                                DemoLead.company.ilike(like), DemoLead.sheet_row_key.ilike(like)))
    rows = (await db.execute(query.order_by(DemoLead.created_at.desc()).limit(limit))).all()
    return [_lead_out(lead, name) for lead, name in rows]


@router.get("/leads/status")
async def leads_status(check: bool = False, p: Principal = Depends(require_demo_admin)):
    """Where leads go, and how the last runs went. check=true also opens the sheet."""
    from app.demo.leads import sheets
    from app.demo.leads.service import STATUS, sheet_on
    from app.demo.settings import csv_list
    s = get_settings()
    out = {"sheet_on": sheet_on(), "sheet_url": sheets.sheet_url(s.demo_leads_gsheet_id.strip()) if sheet_on() else None,
           "tabs": {"leads": s.demo_sheet_leads_tab, "activity": s.demo_sheet_activity_tab or None},
           "email_provider": s.demo_email_provider, "notify_emails": csv_list(s.demo_lead_notify_emails),
           "sync_minutes": s.demo_sheet_sync_minutes, "retry_minutes": s.demo_notify_retry_minutes, **STATUS}
    if check and sheet_on():
        def probe():
            backend = sheets.open_backend()
            return {"title": backend.title, "tabs": backend.tabs()}
        try:
            out["check"] = {"ok": True, **(await sheets.run(probe))}
        except Exception as e:  # noqa: BLE001
            out["check"] = {"ok": False, "error": f"{type(e).__name__}: {str(e)[:300]}"}
    return out


@router.post("/leads/sync")
async def sync_leads_now(p: Principal = Depends(require_demo_admin), db: AsyncSession = Depends(get_db)):
    """Do now what the worker does on its timers: deliver pending leads, read the team's
    columns back, rewrite the Demo activity tab."""
    from app.demo.leads.service import push_pending, sync_sheet
    pushed = await push_pending(db)
    await db.commit()
    synced = await sync_sheet(db)
    await db.commit()
    return {"pushed": pushed, "synced": synced}


@router.post("/leads/{lead_id}/resend")
async def resend_lead(lead_id: uuid.UUID, p: Principal = Depends(require_demo_admin),
                      db: AsyncSession = Depends(get_db)):
    """Retry a lead that failed: its attempts start over and it is delivered now."""
    from app.db.models import DemoLead
    from app.demo.leads.service import push_pending
    lead = await db.get(DemoLead, lead_id)
    if lead is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "lead not found")
    lead.notify_attempts, lead.last_notify_error = 0, None
    append_activity_event(db, p, "demo.lead_resend", resource_type="demo_lead", resource_id=str(lead.id))
    await db.commit()
    result = await push_pending(db)
    await db.commit()
    await db.refresh(lead)
    return {"result": result, "lead": _lead_out(lead, None)}


@router.post("/leads/resync-sheet")
async def resync_sheet(p: Principal = Depends(require_demo_admin), db: AsyncSession = Depends(get_db)):
    """Put back any lead missing from the sheet (rows deleted by mistake). Rows already there
    are left exactly as the team has them."""
    from sqlalchemy import update

    from app.db.models import DemoLead
    from app.demo.leads.service import push_pending, sheet_on
    if not sheet_on():
        raise HTTPException(status.HTTP_409_CONFLICT, "the leads sheet is off (DEMO_LEADS_GSHEET_ID is empty)")
    await db.execute(update(DemoLead).values(sheet_synced_at=None, notify_attempts=0))
    append_activity_event(db, p, "demo.leads_resync_sheet", resource_type="demo_lead")
    await db.commit()
    result = await push_pending(db)
    await db.commit()
    return result


@router.get("/leads.xlsx")
async def leads_xlsx(p: Principal = Depends(require_demo_admin), db: AsyncSession = Depends(get_db)):
    """Every lead as a spreadsheet, built from the database: the fallback if the sheet is
    unavailable."""
    import openpyxl

    from app.db.models import DemoLead
    from app.demo.leads.service import lead_row
    from app.demo.leads.sheets import LEAD_COLUMNS, safe_cell
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Leads"
    ws.append(LEAD_COLUMNS)
    for lead in (await db.execute(select(DemoLead).order_by(DemoLead.created_at))).scalars():
        row = await lead_row(db, lead)
        ws.append([safe_cell(row.get(c)) for c in LEAD_COLUMNS])
    buf = io.BytesIO()
    wb.save(buf)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M")
    return StreamingResponse(io.BytesIO(buf.getvalue()),
                             media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="demo-leads-{stamp}.xlsx"',
                                      "Cache-Control": "no-store"})

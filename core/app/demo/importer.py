"""Read an industry tab of the prospect workbook into company records + contacts
(docs/gtm-demo-links.md §5).

This is the dry run: it never writes. It returns an ImportReport saying, row by row, what
would be imported, what is below the go-live minimum and why, which people need review,
and every problem found. The upload API (phase 2) persists from the same report, so what
the admin previews is exactly what gets saved.

The workbook layout is driven by the industry definition's `sheet` block (tab names, header
markers, column names), so a new industry's tab needs a definition, not code.

Sheets change between imports: the research team adds a column, or renames one. A column
the definition doesn't map yet is not dropped silently. It comes back as a proposal:

  * map  — the header names a fact the definition already has but no column feeds yet
           ("Areas served" → service_area)
  * add  — a new text fact, keyed and labelled from the header

and its values are already in the preview's records. Applying the import with the proposals
accepted publishes them as the industry's next definition version (records.apply_import),
so the new column becomes a normal fact: in the knowledge base, the record screen, the
facts template and every later import. Bookkeeping columns (DEMO_IMPORT_IGNORE_COLUMNS,
sheet.ignore_columns) never become facts.
"""
import difflib
import re
import unicodedata
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, time
from io import BytesIO
from pathlib import Path
from typing import Any, BinaryIO
from zoneinfo import ZoneInfo

import openpyxl
from openpyxl.utils import get_column_letter

from app.config import get_settings
from app.demo.contacts import ParsedPeople, parse_people
from app.demo.facts import FactValue, check_minimum, clean_text, normalize_fact, normalize_url
from app.demo.industries.definition import (DefinitionError, IndustryDefinition,
                                            compile_definition)
from app.demo.sample.validate import validate_dataset
from app.demo.settings import csv_list

_HEADER_SCAN_ROWS = 30
_URL_RE = re.compile(r"https?://\S+|(?:www\.)?[a-z0-9-]+(?:\.[a-z0-9-]+)+\.[a-z]{2,}(?:/\S*)?", re.I)
# An optional "Label:" then a URL. The label may not contain ":" "/" or a newline, so the
# "https:" inside a URL is never read as a label.
_LABELLED_SOURCE_RE = re.compile(r"(?:([A-Za-z][^:/\n;,]{0,59}?)\s*:\s*)?(https?://[^\s,;]+)", re.I)
_TRUE = {"yes", "y", "true", "1", "verified"}
_SAMPLE_CHARS = 80


class WorkbookError(ValueError):
    """The workbook can't be read for this industry at all (no tab, no header row)."""


@dataclass
class RecordPreview:
    row: int
    record_key: str
    company_name: str
    record_name: str | None
    city: str | None
    website_url: str | None
    trigger_note: str | None
    research_verified: bool
    research_status: str | None
    facts: dict[str, FactValue]
    contacts: list[dict]
    company_emails: list[str]
    minimum_missing: list[str]
    dataset_errors: list[str]
    warnings: list[str] = field(default_factory=list)

    @property
    def status(self) -> str:
        # A blank required fact also breaks the placeholders that use it, so "below the
        # minimum" is the root cause to report; "invalid" is for data problems beyond that.
        if self.minimum_missing:
            return "below_minimum"
        return "invalid" if self.dataset_errors else "ready"


@dataclass
class ImportReport:
    industry_key: str
    industry_version: int
    source_name: str
    tab: str
    header_row: int
    records: list[RecordPreview]
    skipped: list[dict]                 # {row, company, reason}
    errors: list[str]                   # sheet-level problems
    warnings: list[str]
    # Columns the definition doesn't map yet (see the module docstring). Their values are
    # in `records` already; accepting them on apply makes them part of the definition.
    new_columns: list[dict] = field(default_factory=list)
    ignored_columns: list[str] = field(default_factory=list)

    @property
    def summary(self) -> dict:
        contacts = [c for r in self.records for c in r.contacts]
        return {
            "records": len(self.records),
            "ready": sum(r.status == "ready" for r in self.records),
            "below_minimum": sum(r.status == "below_minimum" for r in self.records),
            "invalid": sum(r.status == "invalid" for r in self.records),
            "skipped": len(self.skipped),
            "contacts": len(contacts),
            "contacts_with_email": sum(bool(c["email"]) for c in contacts),
            "contacts_needing_review": sum(c["needs_review"] for c in contacts),
            "new_columns": len(self.new_columns),
        }

    def as_json(self) -> dict:
        def fact(fv: FactValue) -> dict:
            d = asdict(fv)
            d["fetched_at"] = fv.fetched_at.isoformat() if fv.fetched_at else None
            return d
        return {
            "industry_key": self.industry_key, "industry_version": self.industry_version,
            "source_name": self.source_name, "tab": self.tab, "header_row": self.header_row,
            "summary": self.summary, "errors": self.errors, "warnings": self.warnings,
            "skipped": self.skipped, "new_columns": self.new_columns,
            "ignored_columns": self.ignored_columns,
            "records": [{**{k: v for k, v in asdict(r).items() if k != "facts"},
                         "status": r.status,
                         "facts": {k: fact(v) for k, v in r.facts.items()}}
                        for r in self.records],
        }


def slugify(text: str) -> str:
    ascii_text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", ascii_text.lower()).strip("-")


def make_record_key(company: str, record_name: str | None) -> str:
    key = slugify(company)
    if record_name:
        key += "--" + slugify(record_name)
    return key[:160]


def norm_header(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().lower()


def _loose(value: str) -> str:
    """Letters and digits only: "Areas served" ~ "areas_served" ~ "Areas-served"."""
    return re.sub(r"[^a-z0-9]+", "", norm_header(value))


def fact_key_for(header: str, taken: set[str]) -> str:
    """A fact key from a column header ("Shipping charges (₹)" → shipping_charges), unique
    against `taken`, matching the definition's key pattern."""
    base = slugify(header).replace("-", "_")[:36].strip("_") or "column"
    if not base[0].isalpha():
        base = f"f_{base}"[:36]
    if len(base) < 2:
        base = f"{base}_x"
    key, n = base, 2
    while key in taken:
        key, n = f"{base}_{n}", n + 1
    return key


def ignored_headers(definition: IndustryDefinition, extra: list[str]) -> set[str]:
    return {norm_header(c) for c in [*extra, *(definition.sheet or {}).get("ignore_columns", [])]}


def find_tab(wb, names: list[str]):
    by_lower = {ws.title.strip().lower(): ws for ws in wb.worksheets}
    for name in names:
        if name.strip().lower() in by_lower:
            return by_lower[name.strip().lower()]
    raise WorkbookError(f"no tab named {' or '.join(map(repr, names))} "
                        f"(tabs: {', '.join(ws.title for ws in wb.worksheets)})")


def find_header(rows: list[tuple], markers: list[str]) -> int:
    want = {norm_header(m) for m in markers}
    for i, row in enumerate(rows[:_HEADER_SCAN_ROWS]):
        if want <= {norm_header(v) for v in row if v is not None}:
            return i
    raise WorkbookError(f"no header row containing {', '.join(markers)} in the first "
                        f"{_HEADER_SCAN_ROWS} rows")


def open_workbook(source: str | Path | bytes | BinaryIO, source_name: str | None = None):
    """(workbook, name). The caller closes the workbook."""
    if isinstance(source, (bytes, bytearray)):
        source = BytesIO(source)
    name = source_name or (Path(source).name if isinstance(source, (str, Path)) else "upload.xlsx")
    try:
        return openpyxl.load_workbook(source, data_only=True, read_only=True), name
    except Exception as e:   # noqa: BLE001 — any unreadable file is one clear error
        raise WorkbookError(f"{name} is not a readable .xlsx workbook ({type(e).__name__})") from None


def read_rows(ws) -> list[tuple]:
    return [tuple(r) for r in ws.iter_rows(values_only=True)]


def _parse_sources(cell: Any, facts_fields: list[dict]) -> tuple[str | None, dict[str, str]]:
    """'Source URLs' → (default URL, {fact_key: URL}). Accepts bare URLs and the
    workbook's labelled form ("Price: https://…", naming the field)."""
    text = clean_text(cell) or ""
    by_label = {}
    for f in facts_fields:
        by_label[f["label"].lower()] = f["key"]
        by_label[f["key"].lower()] = f["key"]
    default, specific = None, {}
    for m in _LABELLED_SOURCE_RE.finditer(text):
        label, url = (m.group(1) or "").strip().lower(), m.group(2).rstrip(".,;)")
        if label in by_label:
            specific.setdefault(by_label[label], url)
        elif default is None:
            default = normalize_url(url)
    if default is None and not specific:
        bare = _URL_RE.search(text)          # "casagrand.co.in" — no scheme, no label
        default = normalize_url(bare.group(0)) if bare else None
    return default, specific


def _fetched_at(cell: Any, tz: ZoneInfo) -> datetime | None:
    if isinstance(cell, datetime):
        return cell.replace(tzinfo=tz) if cell.tzinfo is None else cell
    if isinstance(cell, date):
        return datetime.combine(cell, time(0), tzinfo=tz)
    return None


# ── columns the definition doesn't map yet ───────────────────────────────────────────────

def _match_fact(header: str, candidates: list[dict]) -> dict | None:
    """The fact a new header most likely feeds: same name as its label, key or old column
    (letters and digits only), else a close spelling of its old column (a rename)."""
    loose = _loose(header)
    for f in candidates:
        if loose in {_loose(f["label"]), _loose(f["key"]), _loose(f.get("sheet_column") or "")}:
            return f
    old = {norm_header(f["sheet_column"]): f for f in candidates if f.get("sheet_column")}
    close = difflib.get_close_matches(norm_header(header), list(old), n=1, cutoff=0.8)
    return old[close[0]] if close else None


def propose_columns(definition: IndustryDefinition, header: tuple, data_rows: list[tuple],
                    *, ignore: set[str]) -> tuple[list[dict], list[str]]:
    """(proposals, ignored): one proposal per header the definition doesn't map and that
    isn't in `ignore` (normalised headers), in sheet order.

    A proposal is {header, column, action, key, label, filled, samples[, previous_column]}:
      action "map" — the column feeds an existing fact: one no column fed yet, or one whose
                     column is gone from this sheet (a rename; previous_column says which)
      action "add" — a new optional text fact keyed from the header
    `filled` counts the rows under the header with a value in that column."""
    sheet = definition.sheet or {}
    mapped = {norm_header(v) for v in [*sheet.get("company", {}).values(),
                                        *(sheet.get("contacts") or {}).values()]}
    mapped |= {norm_header(f["sheet_column"]) for f in definition.facts_fields if f.get("sheet_column")}
    present = {norm_header(v) for v in header if clean_text(v)}
    # Facts a new column may feed: no column yet, or their column isn't in this sheet.
    open_facts = [f for f in definition.facts_fields
                  if not f.get("sheet_column") or norm_header(f["sheet_column"]) not in present]
    taken = set(definition.fact_keys)
    proposals, ignored, seen = [], [], set()
    for i, raw_header in enumerate(header):
        text = clean_text(raw_header)
        h = norm_header(text)
        if not text or h in mapped or h in seen:
            continue
        seen.add(h)
        if h in ignore:
            ignored.append(text)
            continue
        values = [v for v in (clean_text(r[i]) for r in data_rows if i < len(r)) if v]
        target = _match_fact(text, open_facts)
        if target is not None:
            open_facts.remove(target)
            proposal = {"action": "map", "key": target["key"], "label": target["label"]}
            if target.get("sheet_column"):
                proposal["previous_column"] = target["sheet_column"]
        else:
            key = fact_key_for(text, taken)
            taken.add(key)
            proposal = {"action": "add", "key": key, "label": text[:120]}
        proposals.append({"header": text, "column": get_column_letter(i + 1), **proposal,
                          "filled": len(values), "samples": [v[:_SAMPLE_CHARS] for v in values[:2]]})
    return proposals, ignored


def apply_column_proposals(raw: dict, proposals: list[dict]) -> dict:
    """`raw` with the proposals in it: a mapped column feeds its fact; an added column is a
    new, optional text fact. The caller compiles (and versions) the result."""
    import copy
    out = copy.deepcopy(raw)
    facts = out["facts_fields"]
    by_key = {f["key"]: f for f in facts}
    for p in proposals:
        if p["action"] == "map" and p["key"] in by_key:
            by_key[p["key"]]["sheet_column"] = p["header"]
        elif p["action"] == "add" and p["key"] not in by_key:
            facts.append({"key": p["key"], "label": p["label"], "type": "text",
                          "sheet_column": p["header"],
                          "description": f"Added from the workbook column {p['header']!r}"})
    return out


def _rename_hints(missing: list[str], candidates: list[str]) -> dict[str, str]:
    """{missing header: the unmapped header it was most likely renamed to}."""
    hints = {}
    for m in missing:
        close = difflib.get_close_matches(norm_header(m), [norm_header(c) for c in candidates],
                                          n=1, cutoff=0.75)
        if close:
            hints[m] = next(c for c in candidates if norm_header(c) == close[0])
    return hints


# ── the dry run ──────────────────────────────────────────────────────────────────────────

def parse_workbook(source: str | Path | bytes | BinaryIO, definition: IndustryDefinition, *,
                   source_name: str | None = None, timezone: str = "Asia/Kolkata",
                   tab: str | None = None, new_columns: str | None = None,
                   ignore_columns: list[str] | None = None,
                   skip_columns: list[str] | None = None) -> ImportReport:
    """Dry-run import of the definition's tab (or `tab`, when the admin picked one) from an
    .xlsx workbook.

    new_columns: "propose" or "ignore" — what to do with columns the definition doesn't map
                 (default DEMO_IMPORT_NEW_COLUMNS); the definition's sheet.extra_columns wins.
    ignore_columns: bookkeeping headers that are never facts (default
                 DEMO_IMPORT_IGNORE_COLUMNS).
    skip_columns: new columns the admin unticked in an earlier preview of this file; they
                 are left out (and listed under ignored_columns) instead of proposed.
    """
    if not definition.sheet:
        raise WorkbookError(f"industry {definition.key!r} has no sheet mapping")
    wb, name = open_workbook(source, source_name)
    try:
        sheet = definition.sheet
        ws = find_tab(wb, [tab] if tab else sheet["tab_names"])
        rows = read_rows(ws)
    finally:
        wb.close()
    tz = ZoneInfo(timezone)
    s = get_settings()
    if ignore_columns is None:
        ignore_columns = csv_list(s.demo_import_ignore_columns)
    new_columns = new_columns or s.demo_import_new_columns
    ignore = ignored_headers(definition, [*(ignore_columns or []), *(skip_columns or [])])

    h = find_header(rows, sheet["header_markers"])
    header = rows[h]
    errors: list[str] = []
    warnings: list[str] = []

    # Columns the definition doesn't know: propose them, and read their values with the
    # definition as it would be once they're accepted, so the preview shows them.
    mode = sheet.get("extra_columns") or new_columns
    proposals, ignored = propose_columns(definition, header, rows[h + 1:], ignore=ignore)
    if proposals and mode == "propose":
        try:
            definition = compile_definition(apply_column_proposals(definition.raw, proposals))
        except DefinitionError as e:          # should not happen; never block the import on it
            warnings.append("new columns could not be added to the definition: " + "; ".join(e.errors))
            proposals = []
    elif proposals:
        warnings.append("columns not in the industry definition, not imported: "
                        + ", ".join(repr(p["header"]) for p in proposals))
        proposals = []

    col = {norm_header(v): i for i, v in enumerate(header) if v is not None and str(v).strip()}

    def idx(name_: str | None) -> int | None:
        return col.get(norm_header(name_)) if name_ else None

    company_cols = {k: idx(v) for k, v in sheet["company"].items()}
    contact_cols = {k: idx(v) for k, v in (sheet.get("contacts") or {}).items()}
    fact_cols = {f["key"]: idx(f.get("sheet_column")) for f in definition.facts_fields}
    if company_cols["company_name"] is None:
        raise WorkbookError(f"column {sheet['company']['company_name']!r} not found")

    unmapped = [clean_text(v) for v in header if clean_text(v)
                and norm_header(v) not in {norm_header(c) for c in [
                    *sheet["company"].values(), *(sheet.get("contacts") or {}).values(),
                    *(f["sheet_column"] for f in definition.facts_fields if f.get("sheet_column"))]}]
    missing_facts = [f for f in definition.facts_fields
                     if f.get("sheet_column") and fact_cols[f["key"]] is None]
    missing_other = [hd for k, hd in {**sheet["company"], **(sheet.get("contacts") or {})}.items()
                     if idx(hd) is None and k != "company_name"]
    hints = _rename_hints([f["sheet_column"] for f in missing_facts] + missing_other, unmapped)

    def renamed(header_: str) -> str:
        return f"; renamed to {hints[header_]!r}?" if header_ in hints else ""

    for f in missing_facts:
        (errors if f.get("required") else warnings).append(
            f"column {f['sheet_column']!r} ({f['label']}) not found{renamed(f['sheet_column'])}")
    for hd in missing_other:
        warnings.append(f"column {hd!r} not found{renamed(hd)}")

    def cell(row: tuple, i: int | None) -> Any:
        return row[i] if i is not None and i < len(row) else None

    example_markers = {m.lower() for m in sheet.get("example_markers", [])}
    skip_statuses = {s.lower() for s in sheet.get("skip_statuses", [])}
    record_name_field = sheet.get("record_name_field")
    records: list[RecordPreview] = []
    skipped: list[dict] = []
    seen_keys: dict[str, int] = {}

    for offset, row in enumerate(rows[h + 1:], start=h + 2):
        if all(clean_text(v) is None for v in row):
            continue
        company = clean_text(cell(row, company_cols["company_name"]))
        if any(clean_text(v) and clean_text(v).lower() in example_markers for v in row[:2]):
            skipped.append({"row": offset, "company": company, "reason": "example row"})
            continue
        if not company:
            skipped.append({"row": offset, "company": None, "reason": "no company name"})
            continue
        status = clean_text(cell(row, company_cols.get("research_status")))
        if status and status.lower() in skip_statuses:
            skipped.append({"row": offset, "company": company, "reason": f"status is {status!r}"})
            continue

        default_src, specific_src = _parse_sources(cell(row, company_cols.get("source_urls")),
                                                   definition.facts_fields)
        website = normalize_url(clean_text(cell(row, company_cols.get("website_url"))) or "")
        fetched = _fetched_at(cell(row, company_cols.get("fetched_at")), tz)
        row_warnings: list[str] = []
        facts: dict[str, FactValue] = {}
        for f in definition.facts_fields:
            fv = normalize_fact(cell(row, fact_cols[f["key"]]), f.get("type", "text"))
            if not fv.is_blank:
                fv.source_url = specific_src.get(f["key"]) or default_src or website
                fv.fetched_at = fetched
            row_warnings += [f"{f['label']}: {w}" for w in fv.warnings]
            facts[f["key"]] = fv

        record_name = facts[record_name_field].value if record_name_field else None
        key = make_record_key(company, record_name)
        if key in seen_keys:
            skipped.append({"row": offset, "company": company,
                            "reason": f"duplicate of row {seen_keys[key]} (same company and "
                                      f"{record_name_field or 'name'})"})
            continue
        seen_keys[key] = offset

        people: ParsedPeople = parse_people(*(clean_text(cell(row, contact_cols.get(k)))
                                              for k in ("full_name", "role", "email", "linkedin_url")))
        row_warnings += people.warnings
        if not people.contacts:
            row_warnings.append("no people listed; add contacts before generating links")

        dataset_errors: list[str] = []
        if definition.has_agent:
            fact_values = {k: v.value for k, v in facts.items()}
            dataset_errors = [str(e) for e in validate_dataset(
                definition.sample_template, definition.sample_schema,
                allowed_placeholders=definition.fact_keys, facts=fact_values)]

        verified = (clean_text(cell(row, company_cols.get("research_verified"))) or "").lower() in _TRUE
        records.append(RecordPreview(
            row=offset, record_key=key, company_name=company, record_name=record_name,
            city=clean_text(cell(row, company_cols.get("city"))), website_url=website,
            trigger_note=clean_text(cell(row, company_cols.get("trigger_note"))),
            research_verified=verified, research_status=status, facts=facts,
            contacts=[{"full_name": c.full_name, "first_name": c.first_name, "role": c.role,
                       "email": c.email, "linkedin_url": c.linkedin_url,
                       "needs_review": c.needs_review, "review_notes": c.notes}
                      for c in people.contacts],
            company_emails=people.company_emails,
            minimum_missing=check_minimum(definition.facts_fields, definition.minimum, facts),
            dataset_errors=dataset_errors, warnings=row_warnings))

    if not records and not errors:
        errors.append("no company rows found under the header")
    return ImportReport(industry_key=definition.key, industry_version=definition.version,
                        source_name=name, tab=ws.title, header_row=h + 1, records=records,
                        skipped=skipped, errors=errors, warnings=warnings,
                        new_columns=proposals, ignored_columns=ignored)

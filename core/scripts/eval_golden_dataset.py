#!/usr/bin/env python3
"""Golden dataset evaluator for DSR Highland Greenz voice agent.

Reads Golden_SingleTurn and Golden_Threads from the Excel golden dataset,
fires each turn against the live API, scores required/forbidden facts, and
writes an Excel report with latency/accuracy breakdown by:
  - language (en / hi / hinglish)
  - direction (inbound / outbound)
  - tool usage (with_tools / without_tools)

Usage (from project root, Docker running):
    python core/scripts/eval_golden_dataset.py \\
        --dataset  "path/to/highland_greenz_golden_dataset.xlsx" \\
        --out      "eval_harness/results/golden_eval_report.xlsx"
"""
import argparse
import json
import re
import sys
import time
import urllib.error
import urllib.request
from collections import defaultdict
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    print("openpyxl not found. Install it: pip install openpyxl")
    sys.exit(1)

BASE = "http://127.0.0.1:8000"
EMAIL = "rahul@advora.ai"
PASSWORD = "rahul12345"
TOKEN = None

GREEN, RED, YELLOW, DIM, OFF = "\033[32m", "\033[31m", "\033[33m", "\033[2m", "\033[0m"

import sys as _sys
if _sys.stdout.encoding and _sys.stdout.encoding.lower() not in ("utf-8", "utf8"):
    import io
    _sys.stdout = io.TextIOWrapper(_sys.stdout.buffer, encoding="utf-8", errors="replace")
    _sys.stderr = io.TextIOWrapper(_sys.stderr.buffer, encoding="utf-8", errors="replace")


# ─────────────────────────────────────── HTTP helpers

def call(method, path, payload=None, timeout=120):
    for attempt in range(2):
        req = urllib.request.Request(
            BASE + path, method=method,
            data=json.dumps(payload).encode() if payload is not None else None,
        )
        req.add_header("Content-Type", "application/json")
        if TOKEN:
            req.add_header("Authorization", f"Bearer {TOKEN}")
        try:
            with urllib.request.urlopen(req, timeout=timeout) as res:
                body = res.read().decode()
                return json.loads(body) if body else None
        except urllib.error.HTTPError as e:
            if e.code == 401 and attempt == 0:
                login()
                continue
            detail = e.read().decode()[:300]
            raise RuntimeError(f"{method} {path} -> {e.code} {detail}") from e


def login():
    global TOKEN
    TOKEN = call("POST", "/auth/login", {"email": EMAIL, "password": PASSWORD})["access_token"]
    print(f"{DIM}Authenticated as {EMAIL}{OFF}")


# ─────────────────────────────────────── Scoring

def _normalize(text: str) -> str:
    t = (text or "").lower()
    t = re.sub(r"(\d),(\d)", r"\1\2", t)
    t = re.sub(r"\brs\.?\s*", "rs", t)
    t = re.sub(r"\brupees?\b", "rs", t)
    t = re.sub(r"\bcrores?\b", "cr", t)
    t = re.sub(r"\blakhs?\b", "lakh", t)
    t = re.sub(r"\bsq\.?\s*ft\.?\b|\bsqft\b|\bsquare\s+f(?:eet|oot|t)\.?\b", "sft", t)
    t = re.sub(r"\b(approx|approximately|around|circa|roughly|about|indicative|only|just)\b", "", t)
    t = re.sub(r"\s+", " ", t).strip()
    return t


def _extract_numbers(text: str) -> list[str]:
    return re.findall(r"\d[\d.]*", text)


def _fact_hit(fact: str, response: str) -> bool:
    if not fact or not response:
        return False
    resp_norm = _normalize(response)
    for clause in re.split(r";", fact):
        clause = clause.strip()
        if not clause:
            continue
        clause_norm = _normalize(clause)
        nums = _extract_numbers(clause_norm)
        key_nums = [n for n in nums if len(n) >= 3]
        if key_nums:
            resp_nums = _extract_numbers(resp_norm)
            num_hits = sum(1 for n in key_nums if n in resp_nums)
            if num_hits / len(key_nums) >= 0.5:
                return True
        STOP = {"that", "this", "with", "from", "have", "will", "your", "their",
                "same", "rate", "across", "configs", "confirm", "sales", "must",
                "builder", "price", "parking", "standard"}
        words = [w for w in re.findall(r"[a-z0-9]+", clause_norm)
                 if len(w) > 3 and w not in STOP]
        if not words:
            continue
        hit_count = sum(1 for w in words if w in resp_norm)
        if hit_count / len(words) >= 0.55:
            return True
    return False


def score_row(required_facts, forbidden_facts, response: str) -> dict:
    req_list = [f.strip() for f in re.split(r"[;\n]", required_facts or "") if f.strip()]
    forb_list = [f.strip() for f in re.split(r"[;\n]", forbidden_facts or "") if f.strip()]
    hits = [f for f in req_list if _fact_hit(f, response)]
    forb_hits = [f for f in forb_list if _fact_hit(f, response)]
    base_score = int(len(hits) / len(req_list) * 100) if req_list else 100
    penalty = min(len(forb_hits) * 20, base_score)
    return {
        "accuracy_score": max(0, base_score - penalty),
        "facts_hit": "; ".join(hits) if hits else None,
        "forbidden_hit": "; ".join(forb_hits) if forb_hits else None,
    }


# ─────────────────────────────────────── Timing extraction

def extract_timing(turn: dict) -> dict:
    debug = turn.get("debug") or {}
    timing = debug.get("timing") or {}
    return {
        "planner_ms":   timing.get("planner_ms") or timing.get("plan_ms"),
        "retrieval_ms": timing.get("retrieval_ms") or timing.get("rag_ms") or timing.get("ret_ms"),
        "responder_ms": timing.get("responder_ms") or timing.get("resp_ms") or timing.get("llm_ms"),
        "tool_ms":      timing.get("tool_ms"),
        "tool_steps":   timing.get("tool_steps") or 0,
        "planner_calls": timing.get("planner_calls") or 1,
        "total_ms":     timing.get("wall_ms") or timing.get("total_ms"),
        "wall_ms":      timing.get("wall_ms"),
    }


def extract_tools(turn: dict) -> tuple[bool, str | None]:
    """Return (has_tools, tools_csv) from executed_tools list."""
    et = turn.get("executed_tools") or []
    if not et:
        # also check timing.tool_steps
        tool_steps = (turn.get("debug") or {}).get("timing", {}).get("tool_steps", 0)
        if tool_steps:
            return True, f"[{tool_steps} steps]"
        return False, None
    names = [str(t.get("tool") or t.get("name") or t) for t in et]
    return True, ", ".join(names)


# ─────────────────────────────────────── Single-turn runner

def run_single_turn(row: dict, verbose: bool = False) -> dict:
    rid = row["id"]
    utterance = row["user_utterance"] or ""
    direction = row.get("channel") or "inbound"  # original direction for reporting

    result = {
        "row_id": rid,
        "lang": row.get("lang"),
        "difficulty": row.get("difficulty"),
        "slice": row.get("slice"),
        "utterance": utterance,
        "response_text": None,
        "direction": direction,
        "expected_intent": row.get("expected_intent"),
        "expected_tool_calls": row.get("expected_tool_calls"),
        "has_tools": False,
        "tools_used": None,
        "planner_ms": None,
        "retrieval_ms": None,
        "responder_ms": None,
        "tool_ms": None,
        "tool_steps": 0,
        "planner_calls": 1,
        "total_ms": None,
        "wall_ms": None,
        "accuracy_score": None,
        "facts_hit": None,
        "forbidden_hit": None,
        "error": None,
        "timestamp": datetime.utcnow().isoformat(),
    }

    t0 = time.perf_counter()
    try:
        session = call("POST", "/sessions", {"channel": "chat", "direction": "inbound"})
        sid = session["session_id"]

        turn = call("POST", f"/sessions/{sid}/turns", {"text": utterance})
        elapsed_ms = (time.perf_counter() - t0) * 1000

        response_text = turn.get("response_text") or ""
        result["response_text"] = response_text

        timing = extract_timing(turn)
        result.update(timing)
        if not result["total_ms"]:
            result["total_ms"] = round(elapsed_ms, 1)
            result["wall_ms"] = round(elapsed_ms, 1)

        has_tools, tools_csv = extract_tools(turn)
        result["has_tools"] = has_tools
        result["tools_used"] = tools_csv

        scored = score_row(row.get("required_facts"), row.get("forbidden_facts"), response_text)
        result.update(scored)

        try:
            call("POST", f"/sessions/{sid}/end", {"outcome": "eval", "analyse": "skip"})
        except Exception:
            pass

        status = f"{GREEN}PASS{OFF}" if result["accuracy_score"] >= 70 else f"{RED}FAIL{OFF}"
        tool_tag = f" [tools:{tools_csv[:30]}]" if has_tools else ""
        if verbose:
            print(f"  {status} {str(rid):6} [{result['accuracy_score']:3}%]  "
                  f"{utterance[:50]}{tool_tag}")
    except Exception as e:
        result["error"] = str(e)[:200]
        result["accuracy_score"] = 0
        result["total_ms"] = round((time.perf_counter() - t0) * 1000, 1)
        if verbose:
            print(f"  {RED}ERR{OFF} {str(rid):6}  {e}")

    return result


# ─────────────────────────────────────── Thread runner

def run_thread(thread_id: str, turns: list[dict], verbose: bool = False) -> list[dict]:
    results = []

    t_session = time.perf_counter()
    try:
        session = call("POST", "/sessions", {"channel": "chat", "direction": "inbound"})
        sid = session["session_id"]
    except Exception as e:
        for t in turns:
            results.append({
                "thread_id": thread_id,
                "turn_no": t.get("turn_no"),
                "lang": t.get("lang"),
                "archetype": t.get("archetype"),
                "utterance": t.get("utterance"),
                "response_text": None,
                "has_tools": False, "tools_used": None,
                "planner_ms": None, "retrieval_ms": None, "responder_ms": None,
                "tool_ms": None, "tool_steps": 0, "planner_calls": 1,
                "total_ms": None, "wall_ms": None,
                "accuracy_score": 0,
                "facts_hit": None, "forbidden_hit": None,
                "error": str(e)[:200],
                "timestamp": datetime.utcnow().isoformat(),
            })
        return results

    turn_resp = None
    for t in turns:
        utterance = t.get("utterance") or ""
        turn_no = t.get("turn_no")
        res = {
            "thread_id": thread_id,
            "turn_no": turn_no,
            "lang": t.get("lang"),
            "archetype": t.get("archetype"),
            "utterance": utterance,
            "expected_intent": t.get("expected_intent"),
            "expected_tool_calls": t.get("expected_tool_calls"),
            "response_text": None,
            "has_tools": False, "tools_used": None,
            "planner_ms": None, "retrieval_ms": None, "responder_ms": None,
            "tool_ms": None, "tool_steps": 0, "planner_calls": 1,
            "total_ms": None, "wall_ms": None,
            "accuracy_score": None,
            "facts_hit": None, "forbidden_hit": None,
            "error": None,
            "timestamp": datetime.utcnow().isoformat(),
        }

        t0 = time.perf_counter()
        try:
            turn_resp = call("POST", f"/sessions/{sid}/turns", {"text": utterance})
            elapsed_ms = (time.perf_counter() - t0) * 1000

            response_text = turn_resp.get("response_text") or ""
            res["response_text"] = response_text

            timing = extract_timing(turn_resp)
            res.update(timing)
            if not res["total_ms"]:
                res["total_ms"] = round(elapsed_ms, 1)
                res["wall_ms"] = round(elapsed_ms, 1)

            has_tools, tools_csv = extract_tools(turn_resp)
            res["has_tools"] = has_tools
            res["tools_used"] = tools_csv

            scored = score_row(t.get("required_facts"), t.get("forbidden_facts"), response_text)
            res.update(scored)

            status = f"{GREEN}PASS{OFF}" if res["accuracy_score"] >= 70 else f"{RED}FAIL{OFF}"
            tool_tag = f" [tools:{tools_csv[:25]}]" if has_tools else ""
            if verbose:
                print(f"    {status} t{turn_no}  [{res['accuracy_score']:3}%]  "
                      f"{utterance[:45]}{tool_tag}")
        except Exception as e:
            res["error"] = str(e)[:200]
            res["accuracy_score"] = 0
            res["total_ms"] = round((time.perf_counter() - t0) * 1000, 1)
            if verbose:
                print(f"    {RED}ERR{OFF} t{turn_no}  {e}")

        results.append(res)

        if turn_resp and (turn_resp.get("end_call") or turn_resp.get("escalation")):
            break

    try:
        call("POST", f"/sessions/{sid}/end", {"outcome": "eval", "analyse": "skip"})
    except Exception:
        pass

    return results


# ─────────────────────────────────────── Excel helpers

HEADER_FILL = PatternFill("solid", fgColor="1E3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
PASS_FILL   = PatternFill("solid", fgColor="D6F5D6")
FAIL_FILL   = PatternFill("solid", fgColor="FFD6D6")
WARN_FILL   = PatternFill("solid", fgColor="FFF4CC")
TOOL_FILL   = PatternFill("solid", fgColor="E8F4FD")  # light blue = tool-using row


def _style_header(ws, row=1):
    for cell in ws[row]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)


def _auto_width(ws, max_w=60):
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        width = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col_letter].width = min(max_w, max(10, width + 2))


def _score_fill(score):
    if score is None:
        return None
    if score >= 80:
        return PASS_FILL
    if score >= 50:
        return WARN_FILL
    return FAIL_FILL


def _percentile(values, p):
    if not values:
        return None
    s = sorted(values)
    idx = int(len(s) * p / 100)
    return round(s[min(idx, len(s) - 1)], 1)


def _median(values):
    if not values:
        return None
    s = sorted(values)
    n = len(s)
    if n % 2 == 1:
        return round(s[n // 2], 1)
    return round((s[n // 2 - 1] + s[n // 2]) / 2, 1)


def _mean(values):
    if not values:
        return None
    return round(sum(values) / len(values), 1)


def _summary_metrics(rows, filter_fn=None):
    """Compute summary stats for a subset of rows."""
    subset = [r for r in rows if not r.get("error")]
    if filter_fn:
        subset = [r for r in subset if filter_fn(r)]
    all_rows = [r for r in rows if not filter_fn or filter_fn(r)] if filter_fn else rows

    times = [r["total_ms"] for r in subset if r.get("total_ms")]
    accs  = [r["accuracy_score"] for r in subset if r.get("accuracy_score") is not None]
    plans = [r["planner_ms"] for r in subset if r.get("planner_ms")]
    rets  = [r["retrieval_ms"] for r in subset if r.get("retrieval_ms")]
    resps = [r["responder_ms"] for r in subset if r.get("responder_ms")]
    tools = [r["tool_ms"] for r in subset if r.get("tool_ms")]
    errors = sum(1 for r in all_rows if r.get("error"))
    pass_rate = round(sum(1 for a in accs if a >= 70) / len(accs) * 100, 1) if accs else None

    return {
        "n": len(all_rows),
        "mean_total_ms": _mean(times),
        "median_total_ms": _median(times),
        "p95_total_ms": _percentile(times, 95),
        "mean_plan_ms": _mean(plans),
        "mean_ret_ms": _mean(rets),
        "mean_resp_ms": _mean(resps),
        "mean_tool_ms": _mean(tools),
        "mean_accuracy": _mean(accs),
        "pass_rate_pct": pass_rate,
        "errors": errors,
        "with_tools": sum(1 for r in subset if r.get("has_tools")),
    }


# ─────────────────────────────────────── Excel writer

def write_excel(st_results: list[dict], th_results: list[dict], out_path: str):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # ── Summary ──────────────────────────────────────────────────────────────
    ws_sum = wb.create_sheet("Summary")
    sum_headers = [
        "sheet", "segment", "filter", "n",
        "mean_total_ms", "median_total_ms", "p95_total_ms",
        "mean_plan_ms", "mean_ret_ms", "mean_resp_ms", "mean_tool_ms",
        "mean_accuracy", "pass_rate_pct", "with_tools", "errors",
    ]
    ws_sum.append(sum_headers)
    _style_header(ws_sum)

    def _add_summary(sheet_label, rows, segment, filter_fn=None, notes=""):
        m = _summary_metrics(rows, filter_fn)
        ws_sum.append([
            sheet_label, segment, notes,
            m["n"], m["mean_total_ms"], m["median_total_ms"], m["p95_total_ms"],
            m["mean_plan_ms"], m["mean_ret_ms"], m["mean_resp_ms"], m["mean_tool_ms"],
            m["mean_accuracy"], m["pass_rate_pct"], m["with_tools"], m["errors"],
        ])

    # SingleTurn breakdowns
    for label, fn, notes in [
        ("ALL", None, "All single-turn rows"),
        ("inbound",  lambda r: r.get("direction") == "inbound",  "inbound rows"),
        ("outbound", lambda r: r.get("direction") == "outbound", "outbound rows"),
        ("with_tools",    lambda r: r.get("has_tools"),          "turns that called tools"),
        ("without_tools", lambda r: not r.get("has_tools"),      "turns with no tool call"),
        ("en",       lambda r: r.get("lang") == "en",      "English"),
        ("hi",       lambda r: r.get("lang") == "hi",      "Hindi"),
        ("hinglish", lambda r: r.get("lang") == "hinglish","Hinglish"),
    ]:
        _add_summary("SingleTurn", st_results, label, fn, notes)

    # Threads breakdowns
    for label, fn, notes in [
        ("ALL", None, "All thread turns"),
        ("with_tools",    lambda r: r.get("has_tools"),   "thread turns that called tools"),
        ("without_tools", lambda r: not r.get("has_tools"),"thread turns with no tool call"),
        ("en",       lambda r: r.get("lang") == "en",       "English"),
        ("hi",       lambda r: r.get("lang") == "hi",       "Hindi"),
        ("hinglish", lambda r: r.get("lang") == "hinglish", "Hinglish"),
    ]:
        _add_summary("Threads", th_results, label, fn, notes)

    _auto_width(ws_sum)

    # ── Latency_Breakdown ─────────────────────────────────────────────────────
    ws_lat = wb.create_sheet("Latency_Breakdown")
    lat_headers = [
        "segment", "category", "n",
        "mean_total_ms", "median_total_ms", "p95_total_ms",
        "mean_plan_ms", "mean_ret_ms", "mean_resp_ms", "mean_tool_ms",
    ]
    ws_lat.append(lat_headers)
    _style_header(ws_lat)

    def _lat_row(seg, cat, rows, fn=None):
        m = _summary_metrics(rows, fn)
        ws_lat.append([
            seg, cat, m["n"],
            m["mean_total_ms"], m["median_total_ms"], m["p95_total_ms"],
            m["mean_plan_ms"], m["mean_ret_ms"], m["mean_resp_ms"], m["mean_tool_ms"],
        ])

    # Single-turn latency by direction
    _lat_row("SingleTurn", "inbound",       st_results, lambda r: r.get("direction") == "inbound")
    _lat_row("SingleTurn", "outbound",      st_results, lambda r: r.get("direction") == "outbound")
    _lat_row("SingleTurn", "with_tools",    st_results, lambda r: r.get("has_tools"))
    _lat_row("SingleTurn", "without_tools", st_results, lambda r: not r.get("has_tools"))
    _lat_row("SingleTurn", "inbound+tools",    st_results,
             lambda r: r.get("direction") == "inbound" and r.get("has_tools"))
    _lat_row("SingleTurn", "inbound+no_tools", st_results,
             lambda r: r.get("direction") == "inbound" and not r.get("has_tools"))
    _lat_row("SingleTurn", "outbound+tools",    st_results,
             lambda r: r.get("direction") == "outbound" and r.get("has_tools"))
    _lat_row("SingleTurn", "outbound+no_tools", st_results,
             lambda r: r.get("direction") == "outbound" and not r.get("has_tools"))

    # Thread latency by tool usage
    _lat_row("Threads", "with_tools",    th_results, lambda r: r.get("has_tools"))
    _lat_row("Threads", "without_tools", th_results, lambda r: not r.get("has_tools"))

    _auto_width(ws_lat)

    # ── Accuracy_Breakdown ────────────────────────────────────────────────────
    ws_acc = wb.create_sheet("Accuracy_Breakdown")
    acc_headers = [
        "segment", "category", "n", "mean_accuracy", "pass_rate_pct",
        "mean_total_ms", "median_total_ms", "errors",
    ]
    ws_acc.append(acc_headers)
    _style_header(ws_acc)

    def _acc_row(seg, cat, rows, fn=None):
        m = _summary_metrics(rows, fn)
        ws_acc.append([
            seg, cat, m["n"], m["mean_accuracy"], m["pass_rate_pct"],
            m["mean_total_ms"], m["median_total_ms"], m["errors"],
        ])

    _acc_row("SingleTurn", "inbound",       st_results, lambda r: r.get("direction") == "inbound")
    _acc_row("SingleTurn", "outbound",      st_results, lambda r: r.get("direction") == "outbound")
    _acc_row("SingleTurn", "with_tools",    st_results, lambda r: r.get("has_tools"))
    _acc_row("SingleTurn", "without_tools", st_results, lambda r: not r.get("has_tools"))
    _acc_row("SingleTurn", "en",       st_results, lambda r: r.get("lang") == "en")
    _acc_row("SingleTurn", "hi",       st_results, lambda r: r.get("lang") == "hi")
    _acc_row("SingleTurn", "hinglish", st_results, lambda r: r.get("lang") == "hinglish")
    _acc_row("Threads",    "with_tools",    th_results, lambda r: r.get("has_tools"))
    _acc_row("Threads",    "without_tools", th_results, lambda r: not r.get("has_tools"))
    _acc_row("Threads",    "en",       th_results, lambda r: r.get("lang") == "en")
    _acc_row("Threads",    "hi",       th_results, lambda r: r.get("lang") == "hi")
    _acc_row("Threads",    "hinglish", th_results, lambda r: r.get("lang") == "hinglish")

    _auto_width(ws_acc)

    # ── SingleTurn_Results ────────────────────────────────────────────────────
    st_cols = [
        "row_id", "lang", "difficulty", "slice", "utterance",
        "response_text", "direction", "has_tools", "tools_used",
        "expected_intent", "expected_tool_calls",
        "planner_ms", "retrieval_ms", "responder_ms", "tool_ms",
        "tool_steps", "planner_calls", "total_ms", "wall_ms",
        "accuracy_score", "facts_hit", "forbidden_hit", "error", "timestamp",
    ]
    ws_st = wb.create_sheet("SingleTurn_Results")
    ws_st.append(st_cols)
    _style_header(ws_st)
    for r in st_results:
        row = [r.get(c) for c in st_cols]
        ws_st.append(row)
        max_row = ws_st.max_row
        score = r.get("accuracy_score")
        fill = _score_fill(score)
        if fill:
            ws_st.cell(max_row, st_cols.index("accuracy_score") + 1).fill = fill
        if r.get("has_tools"):
            ws_st.cell(max_row, st_cols.index("has_tools") + 1).fill = TOOL_FILL
    _auto_width(ws_st)

    # ── Thread_Results ────────────────────────────────────────────────────────
    th_cols = [
        "thread_id", "turn_no", "lang", "archetype", "utterance",
        "response_text", "has_tools", "tools_used",
        "expected_intent", "expected_tool_calls",
        "planner_ms", "retrieval_ms", "responder_ms", "tool_ms",
        "tool_steps", "planner_calls", "total_ms", "wall_ms",
        "accuracy_score", "facts_hit", "forbidden_hit", "error", "timestamp",
    ]
    ws_th = wb.create_sheet("Thread_Results")
    ws_th.append(th_cols)
    _style_header(ws_th)
    for r in th_results:
        row = [r.get(c) for c in th_cols]
        ws_th.append(row)
        max_row = ws_th.max_row
        score = r.get("accuracy_score")
        fill = _score_fill(score)
        if fill:
            ws_th.cell(max_row, th_cols.index("accuracy_score") + 1).fill = fill
        if r.get("has_tools"):
            ws_th.cell(max_row, th_cols.index("has_tools") + 1).fill = TOOL_FILL
    _auto_width(ws_th)

    # ── Accuracy_by_Intent ────────────────────────────────────────────────────
    intent_groups: dict[str, list[dict]] = defaultdict(list)
    for r in st_results:
        intent_groups[r.get("expected_intent") or "unknown"].append(r)

    ws_intent = wb.create_sheet("Accuracy_by_Intent")
    intent_cols = [
        "intent", "n", "mean_accuracy", "pass_rate_pct",
        "mean_total_ms", "median_total_ms", "p95_total_ms",
        "with_tools", "errors",
    ]
    ws_intent.append(intent_cols)
    _style_header(ws_intent)
    for intent, rows in sorted(intent_groups.items()):
        m = _summary_metrics(rows)
        ws_intent.append([
            intent, m["n"], m["mean_accuracy"], m["pass_rate_pct"],
            m["mean_total_ms"], m["median_total_ms"], m["p95_total_ms"],
            m["with_tools"], m["errors"],
        ])
    _auto_width(ws_intent)

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"\n{GREEN}Saved:{OFF} {out_path}")


# ─────────────────────────────────────── Dataset loader

def load_golden_dataset(path: str) -> tuple[list[dict], list[dict]]:
    wb = openpyxl.load_workbook(path)

    def sheet_to_dicts(sheet_name):
        ws = wb[sheet_name]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        rows = []
        for r in range(2, ws.max_row + 1):
            row = {headers[c - 1]: ws.cell(r, c).value
                   for c in range(1, ws.max_column + 1)}
            if all(v is None for v in row.values()):
                continue
            rows.append(row)
        return rows

    return sheet_to_dicts("Golden_SingleTurn"), sheet_to_dicts("Golden_Threads")


# ─────────────────────────────────────── Console summary

def _print_breakdown(label, rows):
    total = len(rows)
    ok = [r for r in rows if not r.get("error")]
    accs = [r["accuracy_score"] for r in ok if r.get("accuracy_score") is not None]
    times = [r["total_ms"] for r in ok if r.get("total_ms")]
    pass_n = sum(1 for a in accs if a >= 70)
    inb = [r for r in ok if r.get("direction") == "inbound"]
    outb = [r for r in ok if r.get("direction") == "outbound"]
    with_t = [r for r in ok if r.get("has_tools")]
    without_t = [r for r in ok if not r.get("has_tools")]

    print(f"\n{'='*60}")
    print(f"  {label}  ({total} rows, {sum(1 for r in rows if r.get('error'))} errors)")
    print(f"{'='*60}")
    print(f"  accuracy  mean={_mean(accs)}%  pass_rate={round(pass_n/len(accs)*100,1) if accs else 0}%")
    print(f"  latency   mean={_mean(times)}ms  median={_median(times)}ms  p95={_percentile(times, 95)}ms")

    if inb:
        inb_t = [r["total_ms"] for r in inb if r.get("total_ms")]
        inb_a = [r["accuracy_score"] for r in inb if r.get("accuracy_score") is not None]
        print(f"\n  INBOUND   n={len(inb)}  acc={_mean(inb_a)}%  "
              f"mean={_mean(inb_t)}ms  median={_median(inb_t)}ms")
    if outb:
        outb_t = [r["total_ms"] for r in outb if r.get("total_ms")]
        outb_a = [r["accuracy_score"] for r in outb if r.get("accuracy_score") is not None]
        print(f"  OUTBOUND  n={len(outb)}  acc={_mean(outb_a)}%  "
              f"mean={_mean(outb_t)}ms  median={_median(outb_t)}ms")

    if with_t:
        wt_t = [r["total_ms"] for r in with_t if r.get("total_ms")]
        wt_a = [r["accuracy_score"] for r in with_t if r.get("accuracy_score") is not None]
        tool_ms = [r["tool_ms"] for r in with_t if r.get("tool_ms")]
        print(f"\n  WITH TOOLS    n={len(with_t)}  acc={_mean(wt_a)}%  "
              f"mean={_mean(wt_t)}ms  median={_median(wt_t)}ms  tool_ms={_mean(tool_ms)}ms")
    if without_t:
        nt_t = [r["total_ms"] for r in without_t if r.get("total_ms")]
        nt_a = [r["accuracy_score"] for r in without_t if r.get("accuracy_score") is not None]
        print(f"  WITHOUT TOOLS n={len(without_t)}  acc={_mean(nt_a)}%  "
              f"mean={_mean(nt_t)}ms  median={_median(nt_t)}ms")


# ─────────────────────────────────────── Main

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dataset", required=True)
    ap.add_argument("--out", default="eval_harness/results/golden_eval_report.xlsx")
    ap.add_argument("--only-lang", default="")
    ap.add_argument("--max-rows", type=int, default=0)
    ap.add_argument("--max-threads", type=int, default=0)
    ap.add_argument("--single-only", action="store_true")
    ap.add_argument("--threads-only", action="store_true")
    ap.add_argument("--verbose", action="store_true", default=True)
    args = ap.parse_args()

    login()

    print(f"\nLoading golden dataset: {args.dataset}")
    single_turn_rows, threads_raw = load_golden_dataset(args.dataset)

    lang_filter = args.only_lang.strip().lower()
    if lang_filter:
        single_turn_rows = [r for r in single_turn_rows if r.get("lang") == lang_filter]
        threads_raw = [r for r in threads_raw if r.get("lang") == lang_filter]

    if args.max_rows:
        single_turn_rows = single_turn_rows[:args.max_rows]

    thread_map: dict[str, list[dict]] = defaultdict(list)
    for t in threads_raw:
        thread_map[t["thread_id"]].append(t)
    for tid in thread_map:
        thread_map[tid].sort(key=lambda x: x.get("turn_no") or 0)

    thread_ids = sorted(thread_map.keys())
    if args.max_threads:
        thread_ids = thread_ids[:args.max_threads]

    # ── Single-turn
    st_results = []
    if not args.threads_only:
        n = len(single_turn_rows)
        print(f"\n{DIM}Single-turn: {n} rows{OFF}")
        for i, row in enumerate(single_turn_rows, 1):
            if i % 10 == 0 or i == n:
                print(f"  {i}/{n} …")
            st_results.append(run_single_turn(row, verbose=args.verbose))

    # ── Threads
    th_results = []
    if not args.single_only:
        n = len(thread_ids)
        print(f"\n{DIM}Threads: {n} threads{OFF}")
        for i, tid in enumerate(thread_ids, 1):
            turns = thread_map[tid]
            if args.verbose:
                print(f"  [{i}/{n}] {tid}  ({len(turns)} turns)")
            th_results.extend(run_thread(tid, turns, verbose=args.verbose))

    # ── Console breakdown
    if st_results:
        _print_breakdown("SINGLE-TURN", st_results)
    if th_results:
        _print_breakdown("THREADS", th_results)

    # ── Write report
    write_excel(st_results, th_results, args.out)

    all_acc = [r.get("accuracy_score") for r in (st_results + th_results)
               if r.get("accuracy_score") is not None]
    errors = sum(1 for r in (st_results + th_results) if r.get("error"))
    print(f"\n  Overall accuracy: {_mean(all_acc)}%   Errors: {errors}")
    return 0 if errors == 0 else 1


if __name__ == "__main__":
    sys.exit(main())

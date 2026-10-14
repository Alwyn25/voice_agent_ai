"""Workbook dry-run importer (app/demo/importer.py) against a synthetic workbook laid out
like the real prospect workbook: title rows above the header, an Example row, a "Not a fit"
row, a duplicate, an error cell, and mismatched people lists. All names are made up.

    cd core && python -m pytest tests/test_demo_importer.py -q
"""
import sys
from datetime import datetime
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.demo.importer import WorkbookError, make_record_key, parse_workbook  # noqa: E402
from app.demo.industries.registry import load_seed_definitions  # noqa: E402

RE = next(d for d in load_seed_definitions() if d.key == "real_estate")

HEADER = ["#", "Developer", "City", "Official website URL", "Buyer name", "Buyer role", "Email",
          "LinkedIn URL", "Trigger / notes", "Verified", "Status", "Project name", "RERA number",
          "Location", "Landmarks", "Configurations", "Price", "Amenities", "Possession date",
          "Approved banks", "Sales office", "Current offer", "Source URLs", "Date scraped",
          "Minimum met?"]


def row(**v):
    return [v.get(h) for h in HEADER]


def workbook(rows, tab="Real estate") -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = tab
    ws.append(["Demo use case: New-project enquiry and site-visit booking"])
    ws.append(["From general market knowledge."])
    ws.append([])
    ws.append([])
    ws.append(HEADER)
    for r in rows:
        ws.append(r)
    wb.create_sheet("Overview").append(["not this tab"])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


READY = row(**{"#": 1, "Developer": "Acme Homes", "City": "Bengaluru",
               "Official website URL": "acmehomes.in",
               "Buyer name": "Asha Rao, vikram iyer", "Buyer role": "Sales head, Manager",
               "Email": "sales@acmehomes.in, asha.rao@acmehomes.in, vikram@acmehomes.in",
               "Verified": "Yes", "Status": "Reviewed", "Project name": "Acme Green Acres",
               "Location": "Whitefield, Bengaluru", "Configurations": "2 & 3 BHK",
               "Price": "₹1.04 Cr onwards", "Possession date": datetime(2029, 12, 1),
               "Sales office": "MG Road, Bengaluru",
               "Source URLs": "https://acmehomes.in/green\nPrice: https://acmehomes.in/price",
               "Date scraped": datetime(2026, 10, 1)})


def report(rows, **kw):
    return parse_workbook(workbook(rows, **kw), RE, source_name="test.xlsx")


def test_ready_record():
    rep = report([READY])
    assert rep.header_row == 5 and rep.tab == "Real estate"
    r = rep.records[0]
    assert r.status == "ready" and r.row == 6
    assert r.record_key == "acme-homes--acme-green-acres"
    assert r.website_url == "https://acmehomes.in/"
    assert r.research_verified is True
    assert r.facts["possession_date"].value == "December 2029"
    assert r.facts["configurations"].items == ["2 BHK", "3 BHK"]
    # Labelled source wins for its field; the bare URL is the default for the rest.
    assert r.facts["price"].source_url == "https://acmehomes.in/price"
    assert r.facts["location"].source_url == "https://acmehomes.in/green"
    assert r.facts["price"].fetched_at.date().isoformat() == "2026-10-01"
    assert r.facts["rera_number"].is_blank and r.facts["rera_number"].source_url is None
    assert [c["full_name"] for c in r.contacts] == ["Asha Rao", "Vikram Iyer"]
    assert [c["email"] for c in r.contacts] == ["asha.rao@acmehomes.in", "vikram@acmehomes.in"]
    assert r.company_emails == ["sales@acmehomes.in"]
    assert rep.summary == {"records": 1, "ready": 1, "below_minimum": 0, "invalid": 0,
                           "skipped": 0, "contacts": 2, "contacts_with_email": 2,
                           "contacts_needing_review": 0, "new_columns": 0}


def test_below_minimum_with_reasons():
    r = report([row(**{"Developer": "Solo Builders", "Project name": "Solo One",
                       "Location": "Rajarhat", "Configurations": "3 BHK"})]).records[0]
    assert r.status == "below_minimum"
    assert r.minimum_missing == ["Sales office is missing",
                                 "Configurations: needs at least 2 BHK types (found 1: 3 BHK)"]


def test_missing_project_name_is_below_minimum_not_invalid():
    # A blank required fact also breaks the {{project_name}} placeholder; the root cause
    # (below minimum) is what's reported as the status.
    r = report([row(**{"Developer": "No Name Co", "Location": "X", "Configurations": "2 & 3 BHK",
                       "Sales office": "Y"})]).records[0]
    assert r.status == "below_minimum"
    assert "Project name is missing" in r.minimum_missing
    assert any("project_name" in e for e in r.dataset_errors)


def test_skips_examples_not_a_fit_blanks_and_duplicates():
    rows = [row(**{"#": "Example", "Developer": "DSR Group", "Project name": "DSR Highland Greenz"}),
            READY,
            [None] * len(HEADER),
            row(**{"Developer": "Gone Co", "Status": "Not a fit", "Project name": "Gone"}),
            row(**{"Developer": None, "Project name": "Orphan"}),
            READY]
    rep = report(rows)
    assert len(rep.records) == 1
    reasons = sorted(s["reason"] for s in rep.skipped)
    assert reasons == ["duplicate of row 7 (same company and project_name)", "example row",
                       "no company name", "status is 'Not a fit'"]


def test_error_cells_become_blank_with_warning():
    r = report([row(**{**dict(zip(HEADER, READY)), "Price": "#ERROR!"})]).records[0]
    assert r.facts["price"].is_blank
    assert any("spreadsheet error cell" in w for w in r.warnings)


def test_missing_required_column_is_a_sheet_error():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Real estate"
    ws.append([h for h in HEADER if h != "Sales office"])
    ws.append(["1", "Acme Homes"])
    buf = BytesIO()
    wb.save(buf)
    rep = parse_workbook(buf.getvalue(), RE)
    assert "column 'Sales office' (Sales office) not found" in rep.errors


def test_wrong_tab_and_no_header():
    with pytest.raises(WorkbookError, match="no tab named 'Real estate'"):
        report([READY], tab="Sheet9")
    wb = openpyxl.Workbook()
    wb.active.title = "Real estate"
    buf = BytesIO()
    wb.save(buf)
    with pytest.raises(WorkbookError, match="no header row"):
        parse_workbook(buf.getvalue(), RE)


def test_not_an_xlsx():
    with pytest.raises(WorkbookError, match="not a readable .xlsx"):
        parse_workbook(b"this is a csv,not,xlsx", RE)


def test_record_key_is_stable_and_ascii():
    assert make_record_key("Concorde Group", "Concorde Élevé") == "concorde-group--concorde-eleve"
    assert make_record_key("Sunteck Realty", "Sunteck City — 4th Avenue") == \
        "sunteck-realty--sunteck-city-4th-avenue"


def test_report_json_is_serialisable():
    import json
    data = report([READY]).as_json()
    json.dumps(data)
    assert data["records"][0]["status"] == "ready"
    assert data["records"][0]["facts"]["price"]["fetched_at"].startswith("2026-10-01")
